"""docx 后端（decisions ㉑A：交付规范来自 docx 实战规范 §8（sources 目录 v1.5.0）；配色、字号、间距一律取 brand/tokens.json，㉕A）。
只读 common.finish_context 之后的 ctx 与 ctx.doc（docmark-ast.md），外加 render.py prepare_figures 的 ctx.fig_info。
入口：build_docx(ctx, out_path) -> stats；readback(ctx, path, stats) -> (readback, layout_issues)；soffice_check(ctx, path, stats)。
DX 规则见 doc-shared/references/layout.md §13。

版式要点：
- 标题用 Word 内置 Heading 1–4 样式（导航窗格、目录可用），编号写成文字（common.heading_label），每个标题、图题、表题带书签，交叉引用为指向书签的内部链接。
- 内容表：tblLayout autofit + tblW auto + 不写 tcW（§8：列宽交给 Word 按内容自适应）；common.table_widths 只写进 tblGrid 作为初始网格。
  首行 tblHeader（标题行重复），每行 cantSplit。表的种类用表格样式区分（DM Table / DM Callout / DM Grid / DM Control / DM Cover / DM Header / DM Footer），读回按样式识别。
- 图：doc-figures build 产出的 figures/<名>.png（2 倍像素）；缺 PNG 放占位段并记 DX4，本后端不栅格化 SVG。位图按显示尺寸降采样，webp 等 Word 不支持的格式用 ImageMagick 转 PNG。
- 脚注：python-docx 无 API，手工加 word/footnotes.xml 部件（原生脚注）；同一脚注再次引用用 NOTEREF 域（Word 的原生做法）。
- 页眉页脚：1 行 2 列无边框表（tblLayout fixed、tblW dxa = 该节版心宽），不用制表位（§8 踩坑）；首节 titlePg，封面页眉页脚为空；横向节断开链接按横向版心重建。
- 目录：TOC 域，结果预填条目（编号 + 标题 + 点线制表位 + PAGEREF 域）；settings.xml updateFields=true，Word 打开时提示更新域。"""
import base64, datetime, hashlib, html, io, os, re, shutil, subprocess, tempfile, zipfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.image.image import Image as DocxImage
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor

import common
from common import heading_label, tilde, table_widths, is_dense, reviewer_text, STATUS_CN, fill, hl_fill, hl_callout_kind

CT_FOOTNOTES = 'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml'
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'
TABLE_STYLES = {'table': 'DM Table', 'callout': 'DM Callout', 'grid': 'DM Grid', 'control': 'DM Control', 'cover': 'DM Cover', 'header': 'DM Header Table', 'footer': 'DM Footer Table'}
LOGO_NAME = 'dm-logo'
CODE_STYLE = 'DM Code'
FN_TEXT_STYLE, FN_REF_STYLE = 'footnote text', 'footnote reference'
LONG_CODE_LINES = 45
BITMAP_MAX_H_FRACTION = 0.6   # layout.md §7：截图最大高度为版心高的 60%
GRID_GAP_MM = 4

# OOXML 子元素顺序（Word 严格校验）
SEQ_PPR = ('w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore', 'w:framePr', 'w:widowControl', 'w:numPr', 'w:suppressLineNumbers',
           'w:pBdr', 'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE',
           'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
           'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle',
           'w:rPr', 'w:sectPr', 'w:pPrChange')
SEQ_RPR = ('w:rStyle', 'w:rFonts', 'w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:caps', 'w:smallCaps', 'w:strike', 'w:dstrike', 'w:outline', 'w:shadow',
           'w:emboss', 'w:imprint', 'w:noProof', 'w:snapToGrid', 'w:vanish', 'w:webHidden', 'w:color', 'w:spacing', 'w:w', 'w:kern',
           'w:position', 'w:sz', 'w:szCs', 'w:highlight', 'w:u', 'w:effect', 'w:bdr', 'w:shd', 'w:fitText', 'w:vertAlign', 'w:rtl', 'w:cs',
           'w:em', 'w:lang', 'w:eastAsianLayout', 'w:specVanish', 'w:oMath')
SEQ_TBLPR = ('w:tblStyle', 'w:tblpPr', 'w:tblOverlap', 'w:bidiVisual', 'w:tblStyleRowBandSize', 'w:tblStyleColBandSize', 'w:tblW', 'w:jc',
             'w:tblCellSpacing', 'w:tblInd', 'w:tblBorders', 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription')
SEQ_TCPR = ('w:cnfStyle', 'w:tcW', 'w:gridSpan', 'w:hMerge', 'w:vMerge', 'w:tcBorders', 'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection',
            'w:tcFitText', 'w:vAlign', 'w:hideMark')
SEQ_SETTINGS = ('w:updateFields', 'w:hdrShapeDefaults', 'w:footnotePr', 'w:endnotePr', 'w:compat', 'w:docVars', 'w:rsids', 'w:attachedSchema',
                'w:themeFontLang', 'w:clrSchemeMapping', 'w:doNotIncludeSubdocsInStats', 'w:doNotAutoCompressPictures', 'w:forceUpgrade',
                'w:captions', 'w:readModeInkLockDown', 'w:smartTagType', 'w:doNotEmbedSmartTags', 'w:decimalSymbol', 'w:listSeparator')
BORDER_SIDES = ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')
LEAK_RE = re.compile(r'⟪|⟫|<!--|-->|\{#[A-Za-z_][\w-]*:|\[\^[^\]\s]+\]|\[![a-z]+\]')


def mm_tw(mm):
    return int(round(mm / 25.4 * 1440))


def mm_pt(mm):
    return mm * 72 / 25.4


def hexc(s):
    return str(s).lstrip('#').upper()


def set_child(parent, tag, seq, attrs=None):
    """在 parent 下按 OOXML 顺序放一个 tag 元素（已有同名元素先删除），返回新元素。"""
    for old in parent.findall(qn(tag)):
        parent.remove(old)
    el = OxmlElement(tag)
    for k, v in (attrs or {}).items():
        el.set(qn(k), str(v))
    succ = {qn(t) for t in seq[seq.index(tag) + 1:]} if tag in seq else set()
    for i, ch in enumerate(parent):
        if ch.tag in succ:
            parent.insert(i, el)
            return el
    parent.append(el)
    return el


def border_el(side, val='single', sz=4, space=0, color='auto'):
    e = OxmlElement(f'w:{side}')
    e.set(qn('w:val'), val)
    if val not in ('nil', 'none'):
        e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), str(space)); e.set(qn('w:color'), hexc(color))
    return e


def docx_filename(ctx):
    return os.path.splitext(common.pdf_filename(ctx))[0] + '.docx'


CJK_FONT_RE = re.compile(r'PingFang|Hiragino|Heiti|Songti|Kaiti|YaHei|SimSun|SimHei|DengXian|Noto\w*(CJK|SC|TC|JP)|SourceHan|Source\s?Han|WenQuanYi|Droid\w*Fallback|STXihei|STSong|FangSong', re.I)
FONT_DIRS = ('/System/Library/Fonts', '/Library/Fonts', '~/Library/Fonts', '/System/Library/AssetsV2/com_apple_MobileAsset_Font7',
             '/System/Library/AssetsV2/com_apple_MobileAsset_Font8', '/usr/share/fonts', '/usr/local/share/fonts')


def fontconfig_file(exe):
    """headless LibreOffice 只看得到自带字体（macOS 上 PingFang 等系统字体进不了 PDF，中文变成不可见字形）。
    生成一份 fonts.conf 把系统字体目录加进去，经 FONTCONFIG_FILE 传给 soffice。"""
    cache = os.path.join(os.path.expanduser('~'), '.cache', 'doc-render'); os.makedirs(cache, exist_ok=True)
    dirs = []
    app_fonts = os.path.join(os.path.dirname(os.path.dirname(os.path.realpath(exe))), 'Resources', 'fonts')
    for dd in (app_fonts,) + FONT_DIRS:
        dd = os.path.expanduser(dd)
        if os.path.isdir(dd): dirs.append(dd)
    body = ''.join(f'  <dir>{html.escape(x)}</dir>\n' for x in dirs)
    conf = ('<?xml version="1.0"?>\n<!DOCTYPE fontconfig SYSTEM "fonts.dtd">\n<fontconfig>\n' + body
            + f'  <cachedir>{html.escape(os.path.join(cache, "fontconfig"))}</cachedir>\n</fontconfig>\n')
    path = os.path.join(cache, 'soffice-fonts.conf')
    if not os.path.exists(path) or open(path, encoding='utf-8').read() != conf:
        open(path, 'w', encoding='utf-8').write(conf)
    return path


def page_residue(ctx, page_text):
    """去掉页眉页脚模板文字与页码后剩下的正文（空串即空白页）。兼容自定义模板：「第 {page} 页」「- {page} -」「{page} / {pages}」等。"""
    words, punct, has_num = [], set(), False
    for spec in ((ctx.profile or {}).get('header') or {}, (ctx.profile or {}).get('footer') or {}):
        for tpl in spec.values():
            tpl = tpl or ''
            if '{page}' in tpl or '{pages}' in tpl: has_num = True
            t_ = re.sub(r'\{(logo|page|pages)\}', ' ', fill(tpl, ctx.vals))
            words += [x for x in re.split(r'\s+', t_) if x]
            punct |= {ch for ch in t_ if not ch.isalnum() and not ch.isspace()}
    rest = page_text or ''
    for x in sorted(set(words), key=len, reverse=True): rest = rest.replace(x, ' ')
    if has_num:
        rest = re.sub(r'\d+', ' ', rest)
        rest = ''.join(ch for ch in rest if ch not in punct)
    return re.sub(r'\s+', '', rest)


def font_problems(pdf_text, pdffonts_out):
    """(问题列表, 已嵌入字体)。pdffonts_out 为 None 表示没有 pdffonts：一律判失败，不假绿。
    只看 emb=yes 的字体；PDF 文字含中文时必须嵌入中文字体，纯拉丁文不要求。"""
    if pdffonts_out is None:
        return ['缺 pdffonts（poppler），无法核验字体嵌入'], []
    embedded = []
    for ln in pdffonts_out.splitlines()[2:]:
        m = re.match(r'^(\S+)\s+.*?\s(yes|no)\s+(yes|no)\s+(yes|no)\s+\d+\s+\d+\s*$', ln)
        if m and m.group(2) == 'yes': embedded.append(m.group(1).split('+')[-1])
    if re.search(r'[\u4e00-\u9fff]', pdf_text or '') and not any(CJK_FONT_RE.search(n) for n in embedded):
        return [f'PDF 含中文但未嵌入中文字体（已嵌入：{", ".join(sorted(set(embedded))[:5]) or "无"}），中文字形不可见'], embedded
    return [], embedded


def merge_soffice(rb, so):
    """把 LibreOffice 核验并入 readback：DX9 转换（必改）、DX8 空白页（建议）；重算 ok（无必改级失败）与 all_passed（全部通过）。"""
    if so is not None:
        rb['checks'].append({'rule': 'DX9', 'name': 'LibreOffice 转 PDF 核验（转换、文字、横向页、字体嵌入）', 'ok': bool(so.get('ok')), 'severity': '必改',
                             'expected': '转换成功且文字、横向页、中文字体一致', 'actual': so.get('error') or f'{so.get("pages")} 页'})
        blank = so.get('blank_pages') or []
        rb['checks'].append({'rule': 'DX8', 'name': 'LibreOffice 转出的 PDF 无空白页', 'ok': not blank, 'severity': '建议', 'expected': '无', 'actual': blank or '无'})
    rb['soffice'] = so
    rb['ok'] = all(c['ok'] for c in rb['checks'] if c.get('severity', '必改') == '必改')
    rb['all_passed'] = all(c['ok'] for c in rb['checks'])
    return rb


def soffice_bin():
    for c in (os.environ.get('DOC_RENDER_SOFFICE'), shutil.which('soffice'), shutil.which('libreoffice'),
              '/Applications/LibreOffice.app/Contents/MacOS/soffice'):
        if c and os.path.isfile(c) and os.access(c, os.X_OK):
            return c
    return None


class DocxBuilder:
    def __init__(self, ctx):
        self.c = ctx; self.doc = ctx.doc; self.t = ctx.tokens
        self.col = ctx.tokens['color']; self.size = ctx.tokens['size']; self.sp = ctx.tokens['space']
        self.d = Document()
        self.used = set()
        self.pending_break = False
        self.cur_landscape = False
        self.bm_id = 0
        self.bm_done = set()
        self.fn_seq = {}          # 脚注 id -> Word 脚注编号（正文发射顺序）
        self.fn_xml = []
        self.abstract = {}
        self.stats = {'headings': 0, 'tables': 0, 'images': 0, 'missing_images': [], 'footnotes': 0, 'footnote_repeats': 0, 'highlights': 0,
                      'code_blocks': 0, 'callouts': 0, 'sections': 1, 'landscape_sections': 0, 'toc': False}
        self.grids = []           # 每张内容表的 (版心 twips, gridCol twips 列表)
        self.toc_entries = 0
        # 书签名：标题、图、表、行内实体锚点共用一张映射；非法字符替换后撞名或超长时加稳定哈希。
        # 书签层级固定一至三级（layout.md §4）：只有一至三级标题是可见书签；四级标题、图、表、实体锚点只作交叉引用目标，用 _ 开头的隐藏书签
        self.bm = {}
        taken = set()
        levels = {eid: h['level'] for eid, h in zip(ctx.heading_ids, ctx.doc.headings)}
        entities = sorted({v for v in ctx.anchor_elem.values() if v.startswith('ent-')})
        for eid in list(ctx.heading_ids) + list(ctx.fig_ids) + list(ctx.tbl_ids) + entities:
            if eid in self.bm:
                continue
            base = re.sub(r'[^A-Za-z0-9_]', '_', eid)
            if not base[:1].isalpha(): base = 'b_' + base
            if levels.get(eid, 99) > ctx.bookmark_levels: base = '_' + base
            name = base if len(base) <= 40 else base[:31] + '_' + hashlib.sha1(eid.encode()).hexdigest()[:8]
            if name in taken:
                name = base[:31] + '_' + hashlib.sha1(eid.encode()).hexdigest()[:8]
            taken.add(name); self.bm[eid] = name

    # ---------- 样式与页面
    def fonts(self, rpr, mono=False):
        f = self.t['font']
        latin = (f['mono'] if mono else f['latin'])[0]
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = set_child(rpr, 'w:rFonts', SEQ_RPR)
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        rf.set(qn('w:ascii'), latin); rf.set(qn('w:hAnsi'), latin); rf.set(qn('w:eastAsia'), f['cjk'][0]); rf.set(qn('w:cs'), latin)

    def pstyle(self, name, base='Normal', size=None, color=None, bold=None, line_pt=None, line_mult=None, before=0.0, after=0.0, align=None, keep_next=None, mono=False):
        st = self.d.styles
        try:
            s = st[name]
        except KeyError:
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = st[base]
        s.hidden = False
        self.fonts(s.element.get_or_add_rPr(), mono)
        if size: s.font.size = Pt(size)
        if color: s.font.color.rgb = RGBColor.from_string(hexc(color))
        if bold is not None: s.font.bold = bold
        s.font.italic = False
        pf = s.paragraph_format
        if line_pt is not None: pf.line_spacing = Pt(line_pt)
        if line_mult is not None: pf.line_spacing = line_mult
        pf.space_before = Pt(before); pf.space_after = Pt(after)
        if align is not None: pf.alignment = align
        if keep_next is not None: pf.keep_with_next = keep_next
        return s

    def bottom_rule(self, ppr, color, sz=12, space=4):
        bdr = set_child(ppr, 'w:pBdr', SEQ_PPR)
        bdr.append(border_el('bottom', sz=sz, space=space, color=color))

    def setup(self):
        sz = self.size; col = self.col; lh = self.t['line_height']
        dd = self.d.styles.element.find(qn('w:docDefaults'))
        rpr = dd.find(qn('w:rPrDefault') + '/' + qn('w:rPr')) if dd is not None else None
        if rpr is not None:
            self.fonts(rpr)
            set_child(rpr, 'w:lang', SEQ_RPR, {'w:val': 'en-US', 'w:eastAsia': self.c.meta.get('language') or 'zh-CN'})
        self.pstyle('Normal', size=sz['body'], color=col['ink'], line_pt=sz['body'] * lh['body'], after=mm_pt(self.sp['para']))
        for lv, size, color, before, after in ((1, sz['h1'], col['ink'], 9, 4), (2, sz['h2'], col['primary_dark'], 6, 3), (3, sz['h3'], col['ink'], 4, 2), (4, sz['body'], col['ink'], 3, 1.5)):
            s = self.pstyle(f'Heading {lv}', size=size, color=color, bold=True, line_mult=1.2, before=mm_pt(before), after=mm_pt(after), keep_next=True)
            s.paragraph_format.keep_together = True
            if lv == 1: self.bottom_rule(s.element.get_or_add_pPr(), col['primary'])
        self.pstyle('Caption', size=sz['caption'], color=col['ink_2'], bold=False, line_mult=1.3, before=mm_pt(2), after=mm_pt(1.5), keep_next=True)
        self.pstyle('DM Figure Caption', base='Caption', size=sz['caption'], color=col['muted'], line_mult=1.3, before=mm_pt(1.5), after=mm_pt(self.sp['figure']),
                    align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=False)
        self.pstyle('DM Figure', line_mult=1.0, before=mm_pt(4), after=0, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        self.pstyle('DM Table Text', size=sz['table'], line_pt=sz['table'] * lh['table'], after=0)
        self.pstyle('DM Table Text Dense', base='DM Table Text', size=sz['table_dense'], line_pt=sz['table_dense'] * lh['table'], after=0)
        s = self.pstyle(CODE_STYLE, size=sz['code'], color=col['ink'], line_pt=sz['code'] * lh['code'], after=0, mono=True)
        ppr = s.element.get_or_add_pPr()
        bdr = set_child(ppr, 'w:pBdr', SEQ_PPR)
        for side in ('left', 'right'): bdr.append(border_el(side, sz=4, space=8, color=col['surface']))
        set_child(ppr, 'w:shd', SEQ_PPR, {'w:val': 'clear', 'w:color': 'auto', 'w:fill': hexc(col['surface'])})
        s.paragraph_format.left_indent = Mm(3); s.paragraph_format.right_indent = Mm(3)
        s = self.pstyle('DM Front Title', size=sz['h1'], color=col['ink'], bold=True, line_mult=1.2, after=mm_pt(5), keep_next=True)
        self.bottom_rule(s.element.get_or_add_pPr(), col['primary'])
        self.pstyle(FN_TEXT_STYLE, size=sz['footnote'], color=col['ink_2'], line_mult=1.2, after=0)
        try:
            cs = self.d.styles[FN_REF_STYLE]
        except KeyError:
            cs = self.d.styles.add_style(FN_REF_STYLE, WD_STYLE_TYPE.CHARACTER)
        cs.font.superscript = True
        self.pstyle('DM Header', size=sz['header_footer'], color=col['subtle'], line_mult=1.0, after=0)
        for lv, size, color, before, indent in ((1, sz['body'], col['ink'], 2.4, 0), (2, sz['table'], col['ink_2'], 0.6, 9)):
            s = self.pstyle(f'toc {lv}', size=size, color=color, bold=(lv == 1), line_mult=1.3, before=mm_pt(before), after=0)
            s.paragraph_format.left_indent = Mm(indent)
            s.paragraph_format.tab_stops.add_tab_stop(Mm(self.c.content_mm - indent), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        for name in TABLE_STYLES.values():
            try:
                self.d.styles[name]
            except KeyError:
                ts = self.d.styles.add_style(name, WD_STYLE_TYPE.TABLE)
                ts.base_style = self.d.styles['Normal Table']
        self.page(self.d.sections[0], landscape=False)
        pkg = self.d.part.package
        for rid, rel in list(pkg.rels.items()):   # python-docx 默认模板自带与本文无关的缩略图
            if rel.reltype == RT.THUMBNAIL:
                del pkg.rels[rid]

    def page(self, sec, landscape):
        p = self.t['page']; sp = self.sp
        w, h = Mm(p['width_mm']), Mm(p['height_mm'])
        sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = (h, w) if landscape else (w, h)
        sec.top_margin = Mm(sp['page_margin_top']); sec.bottom_margin = Mm(sp['page_margin_bottom'])
        sec.left_margin = Mm(sp['page_margin_x']); sec.right_margin = Mm(sp['page_margin_x'])
        sec.header_distance = Mm(8); sec.footer_distance = Mm(7)

    def limit_mm(self):
        return self.c.landscape_mm if self.cur_landscape else self.c.content_mm

    # ---------- 行内
    def run(self, parent, text=None, f=None):
        f = f or {}
        r = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        if f.get('rstyle'):
            e = OxmlElement('w:rStyle'); e.set(qn('w:val'), f['rstyle']); rpr.append(e)
        if f.get('b'):
            rpr.append(OxmlElement('w:b')); rpr.append(OxmlElement('w:bCs'))
        if f.get('caps'): rpr.append(OxmlElement('w:caps'))
        if f.get('color'):
            e = OxmlElement('w:color'); e.set(qn('w:val'), hexc(f['color'])); rpr.append(e)
        if f.get('size'):
            for tag in ('w:sz', 'w:szCs'):
                e = OxmlElement(tag); e.set(qn('w:val'), str(int(round(f['size'] * 2)))); rpr.append(e)
        if f.get('shd'):
            e = OxmlElement('w:shd'); e.set(qn('w:val'), 'clear'); e.set(qn('w:color'), 'auto'); e.set(qn('w:fill'), hexc(f['shd'])); rpr.append(e)
        if f.get('vert'):
            e = OxmlElement('w:vertAlign'); e.set(qn('w:val'), f['vert']); rpr.append(e)
        if len(rpr): r.append(rpr)
        if text:
            for i, s in enumerate(re.split(r'\r?\n', text)):
                if i: r.append(OxmlElement('w:br'))
                if s:
                    te = OxmlElement('w:t'); te.set(XML_SPACE, 'preserve'); te.text = s; r.append(te)
        parent.append(r)
        return r

    def br(self, parent):
        r = OxmlElement('w:r'); r.append(OxmlElement('w:br')); parent.append(r)

    def inline(self, parent, nodes, f=None, where='body'):
        """where：body（正文、表格、高亮块）/ footnote（脚注正文）/ toc（目录条目）/ cover（封面元信息）。
        只有 body 发射脚注与外链关系；其余位置的脚注引用不输出，外链退化为文字。"""
        f = f or {}
        for n in nodes:
            t = n['t']
            if t == 'text':
                self.run(parent, tilde(self.c, n['v']), f)
            elif t == 'bold':
                self.inline(parent, n['c'], {**f, 'b': True}, where)
            elif t == 'highlight':
                if not f.get('shd') and where in ('body', 'footnote'): self.stats['highlights'] += 1
                self.inline(parent, n['c'], {**f, 'b': True, 'shd': hl_fill(self.t, n)} if where in ('body', 'footnote') else {**f, 'b': True}, where)
            elif t == 'anchor':   # 行内实体锚点：第一处发射隐藏书签，交叉引用链接指向它
                eid = self.c.anchor_elem.get(n['target'])
                name = self.bm.get(eid) if eid and eid.startswith('ent-') else None
                if name and name not in self.bm_done and where in ('body', 'footnote'):
                    self.bm_done.add(name); self.bm_id += 1
                    e = OxmlElement('w:bookmarkStart'); e.set(qn('w:id'), str(self.bm_id)); e.set(qn('w:name'), name); parent.append(e)
                    e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), str(self.bm_id)); parent.append(e)
            elif t == 'link':
                url = n['url']
                inner = n['c'] or [{'t': 'text', 'v': url}]
                if where == 'body' and re.match(r'^(https?:|mailto:)', url):
                    h = OxmlElement('w:hyperlink')
                    h.set(qn('r:id'), self.d.part.relate_to(url, RT.HYPERLINK, is_external=True))
                    self.inline(h, inner, {**f, 'color': f.get('color') or self.col['primary_dark']}, where)
                    parent.append(h)
                else:
                    self.inline(parent, inner, f, where)
            elif t == 'ref':
                if not n.get('resolved'):
                    self.run(parent, '@' + n['target'], {**f, 'b': True, 'color': self.col['danger']})
                    continue
                lab = n.get('label') or n['target']
                eid = self.c.anchor_elem.get(n['target'])
                if eid and eid in self.bm and where in ('body', 'footnote'):
                    h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), self.bm[eid]); h.set(qn('w:history'), '1')
                    self.run(h, lab, {**f, 'color': f.get('color') or self.col['primary_dark']})
                    parent.append(h)
                else:
                    self.run(parent, lab, f)
            elif t == 'fnref':
                if where == 'body':
                    self.footnote_ref(parent, n['id'], f)

    def footnote_ref(self, parent, fid, f):
        fn = self.doc.footnotes.get(fid)
        if not fn or fid not in self.c.fn_num:
            return
        rs = self.d.styles[FN_REF_STYLE].style_id
        if fid in self.fn_seq:   # 再次引用：NOTEREF 域指向第一次引用处的书签
            wid = self.fn_seq[fid]
            self.field(parent, f'NOTEREF _Ref_dmfn{wid} \\f \\h', str(wid), {'rstyle': rs})
            self.stats['footnote_repeats'] += 1
            return
        wid = len(self.fn_seq) + 1
        self.fn_seq[fid] = wid
        self.bm_id += 1
        s = OxmlElement('w:bookmarkStart'); s.set(qn('w:id'), str(self.bm_id)); s.set(qn('w:name'), f'_Ref_dmfn{wid}'); parent.append(s)
        r = self.run(parent, None, {'rstyle': rs})
        e = OxmlElement('w:footnoteReference'); e.set(qn('w:id'), str(wid)); r.append(e)
        e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), str(self.bm_id)); parent.append(e)
        p = OxmlElement('w:p')
        ppr = OxmlElement('w:pPr'); ps = OxmlElement('w:pStyle'); ps.set(qn('w:val'), self.d.styles[FN_TEXT_STYLE].style_id); ppr.append(ps); p.append(ppr)
        self.run(p, None, {'rstyle': rs}).append(OxmlElement('w:footnoteRef'))
        self.run(p, ' ')
        self.inline(p, fn['inline'], {}, where='footnote')
        self.fn_xml.append((wid, p))
        self.stats['footnotes'] += 1

    # ---------- 段落、书签、域
    def new_p(self, cont, style=None):
        if cont is self.d:
            p = self.d.add_paragraph(style=style)
            if self.pending_break:
                p.paragraph_format.page_break_before = True
                self.pending_break = False
        else:
            ps = cont.paragraphs
            if len(ps) == 1 and ps[0]._p not in self.used and not ps[0]._p.findall(qn('w:r')):
                p = ps[0]
                if style: p.style = style
            else:
                p = cont.add_paragraph(style=style)
        self.used.add(p._p)
        return p

    def spacer(self, cont, after_pt, keep_next=False):
        p = self.new_p(cont)
        pf = p.paragraph_format
        pf.space_before = Pt(0); pf.space_after = Pt(max(0.0, after_pt)); pf.line_spacing = Pt(2); pf.keep_with_next = keep_next
        return p

    def flush_break_for_table(self, cont):
        if cont is self.d and self.pending_break:
            self.spacer(cont, 0, keep_next=True)

    def bookmark(self, p, eid):
        name = self.bm.get(eid)
        if not name or name in self.bm_done:
            return
        self.bm_done.add(name)
        self.bm_id += 1
        s = OxmlElement('w:bookmarkStart'); s.set(qn('w:id'), str(self.bm_id)); s.set(qn('w:name'), name)
        e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), str(self.bm_id))
        ppr = p._p.find(qn('w:pPr'))
        if ppr is not None: ppr.addnext(s)
        else: p._p.insert(0, s)
        p._p.append(e)

    def fld_char(self, parent, kind, f=None, dirty=False):
        r = self.run(parent, None, f)
        fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind)
        if dirty: fc.set(qn('w:dirty'), 'true')
        r.append(fc)

    def instr(self, parent, text, f=None):
        r = self.run(parent, None, f)
        it = OxmlElement('w:instrText'); it.set(XML_SPACE, 'preserve'); it.text = f' {text} '; r.append(it)

    def field(self, parent, instr, result='', f=None):
        self.fld_char(parent, 'begin', f); self.instr(parent, instr, f); self.fld_char(parent, 'separate', f)
        if result: self.run(parent, result, f)
        self.fld_char(parent, 'end', f)

    # ---------- 表格公共
    def tbl_props(self, tbl, kind, borders=None, layout='autofit', width_tw=None, cell_mar=None, caption=None):
        tblPr = tbl._tbl.tblPr
        set_child(tblPr, 'w:tblStyle', SEQ_TBLPR, {'w:val': self.d.styles[TABLE_STYLES[kind]].style_id})
        if width_tw is None:
            set_child(tblPr, 'w:tblW', SEQ_TBLPR, {'w:type': 'auto', 'w:w': '0'})
        else:
            set_child(tblPr, 'w:tblW', SEQ_TBLPR, {'w:type': 'dxa', 'w:w': str(width_tw)})
        set_child(tblPr, 'w:tblInd', SEQ_TBLPR, {'w:w': '0', 'w:type': 'dxa'})
        b = set_child(tblPr, 'w:tblBorders', SEQ_TBLPR)
        for side in BORDER_SIDES:
            spec = (borders or {}).get(side)
            b.append(border_el(side, sz=spec[0], color=spec[1]) if spec else border_el(side, 'nil'))
        set_child(tblPr, 'w:tblLayout', SEQ_TBLPR, {'w:type': layout})
        mar = set_child(tblPr, 'w:tblCellMar', SEQ_TBLPR)
        my, mx = cell_mar or (self.sp['cell_y'], self.sp['cell_x'])
        for side, v in (('top', my), ('left', mx), ('bottom', my), ('right', mx)):
            e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(mm_tw(v))); e.set(qn('w:type'), 'dxa'); mar.append(e)
        set_child(tblPr, 'w:tblLook', SEQ_TBLPR, {'w:val': '0000', 'w:firstRow': '0', 'w:lastRow': '0', 'w:firstColumn': '0', 'w:lastColumn': '0', 'w:noHBand': '1', 'w:noVBand': '1'})
        for old in tblPr.findall(qn('w:tblDescription')): tblPr.remove(old)
        if caption: set_child(tblPr, 'w:tblCaption', SEQ_TBLPR, {'w:val': caption[:250]})

    def grid_cols(self, tbl, widths_tw):
        grid = tbl._tbl.tblGrid
        for gc in list(grid): grid.remove(gc)
        for w in widths_tw:
            gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(w))); grid.append(gc)

    def strip_tcw(self, tbl):
        for tc in tbl._tbl.iter(qn('w:tc')):
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is not None:
                for e in tcPr.findall(qn('w:tcW')): tcPr.remove(e)

    def set_tcw(self, cell, tw):
        set_child(cell._tc.get_or_add_tcPr(), 'w:tcW', SEQ_TCPR, {'w:w': str(int(tw)), 'w:type': 'dxa'})

    def cell_shd(self, cell, fill_hex):
        set_child(cell._tc.get_or_add_tcPr(), 'w:shd', SEQ_TCPR, {'w:val': 'clear', 'w:color': 'auto', 'w:fill': hexc(fill_hex)})

    def cell_borders(self, cell, spec):
        tb = set_child(cell._tc.get_or_add_tcPr(), 'w:tcBorders', SEQ_TCPR)
        for side in ('top', 'left', 'bottom', 'right'):
            tb.append(border_el(side, sz=spec[side][0], color=spec[side][1]) if spec.get(side) else border_el(side, 'nil'))

    def valign(self, cell, v='center'):
        set_child(cell._tc.get_or_add_tcPr(), 'w:vAlign', SEQ_TCPR, {'w:val': v})

    def row_flags(self, row, header=False):
        trPr = row._tr.get_or_add_trPr()
        if not trPr.findall(qn('w:cantSplit')): trPr.append(OxmlElement('w:cantSplit'))
        if header and not trPr.findall(qn('w:tblHeader')): trPr.append(OxmlElement('w:tblHeader'))

    # ---------- 块
    def block(self, cont, b):
        k = b['kind']
        if k == 'p':
            p = self.new_p(cont)
            self.inline(p._p, b['inline'])
        elif k in ('ul', 'ol'):
            self.lst(cont, b)
        elif k in ('table', 'data'):
            self.table(cont, b)
        elif k in ('figure', 'image'):
            self.figure(cont, b)
        elif k == 'callout':
            self.callout(cont, b)
        elif k == 'code':
            self.code(cont, b)
        elif k == 'pagebreak':
            self.pending_break = True

    def abstract_num(self, k0, k1):
        key = (k0, k1)
        if key in self.abstract:
            return self.abstract[key]
        numbering = self.d.part.numbering_part.element
        ids = [int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))]
        aid = max(ids + [100]) + 1
        lv = []
        for il, kind in enumerate((k0, k1)):
            fmt, txt = ('decimal', f'%{il + 1}.') if kind == 'ol' else ('bullet', '•' if il == 0 else '◦')
            lv.append(f'<w:lvl w:ilvl="{il}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{txt}"/><w:lvlJc w:val="left"/>'
                      f'<w:pPr><w:ind w:left="{mm_tw(6 * (il + 1))}" w:hanging="{mm_tw(5)}"/></w:pPr></w:lvl>')
        el = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}"><w:multiLevelType w:val="hybridMultilevel"/>{"".join(lv)}</w:abstractNum>')
        first_num = numbering.find(qn('w:num'))
        if first_num is not None: first_num.addprevious(el)
        else: numbering.append(el)
        self.abstract[key] = aid
        return aid

    def lst(self, cont, b):
        items = b['items']
        k1 = next((it['kind'] for it in items if it['level'] == 2), b['kind'])
        num = self.d.part.numbering_part.element.add_num(self.abstract_num(b['kind'], k1))
        for il in (0, 1):
            num.add_lvlOverride(ilvl=il).add_startOverride(1)   # 每个列表从 1 重新编号
        in_cell = cont is not self.d
        for i, it in enumerate(items):
            p = self.new_p(cont)
            np_ = set_child(p._p.get_or_add_pPr(), 'w:numPr', SEQ_PPR)
            e = OxmlElement('w:ilvl'); e.set(qn('w:val'), str(min(1, it['level'] - 1))); np_.append(e)
            e = OxmlElement('w:numId'); e.set(qn('w:val'), str(num.numId)); np_.append(e)
            last = i == len(items) - 1
            p.paragraph_format.space_after = Pt(0 if in_cell and last else (mm_pt(self.sp['para']) if last else mm_pt(0.8)))
            self.inline(p._p, it['inline'])

    def table(self, cont, b):
        c = self.c
        if b.get('error'):
            p = self.new_p(cont)
            self.run(p._p, f'数据块无法渲染：{b.get("data_file") or ""}', {'color': self.col['danger'], 'b': True})
            return
        ti = c.tbl_index.get(id(b))
        eid = c.tbl_ids[ti] if ti is not None else ''
        limit = self.limit_mm()
        pcts, _ = table_widths(c, b)
        tstyle = 'DM Table Text Dense' if is_dense(c, b) else 'DM Table Text'
        cap_text = ' '.join(x for x in (b.get('number'), tilde(c, b.get('caption') or '')) if x)
        if cap_text:
            p = self.new_p(cont, 'Caption')
            if b.get('number'): self.run(p._p, b['number'] + '  ', {'b': True, 'color': self.col['primary_dark']})
            if b.get('caption'): self.run(p._p, tilde(c, b['caption']))
            if eid: self.bookmark(p, eid)
        else:
            self.flush_break_for_table(cont)
            if eid and cont is self.d: self.bookmark(self.spacer(cont, 0, keep_next=True), eid)
        ncol = max(1, len(b['header']))
        starts = {g['start']: g for g in (b.get('groups') or [])}
        tbl = cont.add_table(rows=1 + len(b['rows_inline']) + len(starts), cols=ncol)
        self.tbl_props(tbl, 'table', borders={s: (4, self.col['border']) for s in BORDER_SIDES}, caption=cap_text or None)
        grid_tw = [mm_tw(limit * w / 100) for w in pcts]
        self.grid_cols(tbl, grid_tw)
        self.grids.append((mm_tw(limit), grid_tw))
        rows = tbl.rows
        hdr = rows[0]
        self.row_flags(hdr, header=c.features.get('thead_repeat') is not False)
        for ci, x in enumerate(b['header_inline'][:ncol]):
            cell = hdr.cells[ci]
            self.cell_shd(cell, self.col['surface'])
            self.inline(self.new_p(cell, tstyle)._p, x, {'b': True})
        ro = 1
        for ri, r in enumerate(b['rows_inline']):
            if ri in starts:
                g = starts[ri]
                row = rows[ro]; ro += 1
                self.row_flags(row)
                merged = row.cells[0].merge(row.cells[ncol - 1]) if ncol > 1 else row.cells[0]
                self.cell_shd(merged, self.col['tint'])
                self.run(self.new_p(merged, tstyle)._p, f'{tilde(c, g["value"] or "（空）")}（{g["count"]} 行）', {'b': True, 'color': self.col['primary_dark']})
            row = rows[ro]; ro += 1
            self.row_flags(row)
            cells = row.cells
            for ci in range(ncol):
                cell = cells[ci]
                if ri % 2 == 1: self.cell_shd(cell, self.col['surface_2'])
                p = self.new_p(cell, tstyle)
                if ci < len(r): self.inline(p._p, r[ci])
                raw = b['rows'][ri][ci].strip() if ri < len(b['rows']) and ci < len(b['rows'][ri]) else ''
                if raw and len(raw) <= common.NOWRAP_CODE_MAX and common.RE_CODE_CELL.match(raw):
                    set_child(cell._tc.get_or_add_tcPr(), 'w:noWrap', SEQ_TCPR)
        self.strip_tcw(tbl)   # §8：合并单元格之后统一清掉 tcW，列宽交给 Word 自适应
        self.stats['tables'] += 1
        if cont is self.d:
            self.spacer(cont, mm_pt(self.sp['table']))

    def figure_file(self, b):
        """(路径, 期望的运行目录相对路径)：figure 用 doc-figures 的 figures/<名>.png；image 用 src 位图。"""
        if b['kind'] == 'image':
            rel = b['src']
        else:
            base = b.get('svg') or b['src']
            base = base[:-len('.fig.json')] if base.endswith('.fig.json') else os.path.splitext(base)[0]
            rel = base + '.png'
        path = os.path.join(self.c.run_dir, rel)
        if not common.dp.path_safe(rel) or not common.dp.within(self.c.run_dir, path) or not os.path.isfile(path):
            return None, rel
        return path, rel

    def image_bytes(self, b, path, width_mm):
        if b['kind'] != 'image':
            return open(path, 'rb').read()
        uri, _ = common.image_data_uri(self.c, b['src'], width_mm)
        data = base64.b64decode(uri.split(',', 1)[1]) if uri else open(path, 'rb').read()
        try:
            DocxImage.from_blob(data)
            return data
        except Exception:  # noqa: BLE001  Word 不支持的格式（如 webp）转 PNG
            mg = common.magick()
            if not mg:
                raise
            r = subprocess.run([mg, path + '[0]', 'png:-'] if mg.endswith('magick') else [mg, path + '[0]', 'png:-'], capture_output=True, timeout=60)
            if r.returncode != 0 or not r.stdout:
                raise
            return r.stdout

    def picture(self, p, b, width_mm, max_h_mm):
        path, rel = self.figure_file(b)
        if not path:
            self.stats['missing_images'].append({'src': b['src'], 'expected': rel})
            self.run(p._p, f'图未嵌入：{rel} 不存在（先运行 doc-figures build 产出 PNG）', {'color': self.col['danger'], 'b': True, 'size': self.size['caption']})
            return False
        try:
            data = self.image_bytes(b, path, width_mm)
            img = DocxImage.from_blob(data)
        except Exception as ex:  # noqa: BLE001
            self.stats['missing_images'].append({'src': b['src'], 'expected': rel, 'error': str(ex)[:120]})
            self.run(p._p, f'图片无法嵌入：{rel}', {'color': self.col['danger'], 'b': True, 'size': self.size['caption']})
            return False
        pw, ph = img.px_width or 1, img.px_height or 1
        w = min(width_mm, max_h_mm * pw / ph)
        run = p.add_run()
        run.add_picture(io.BytesIO(data), width=Mm(w))
        docpr = run._r.findall('.//' + qn('wp:docPr'))
        if docpr:
            fi = self.c.fig_index.get(id(b))
            docpr[-1].set('name', f'dm-figure-{(fi if fi is not None else 0) + 1}')
            docpr[-1].set('descr', ' '.join(x for x in (b.get('number'), b.get('caption')) if x)[:250])
        self.stats['images'] += 1
        return True

    def fig_dims(self, b, grid_cols=None):
        c = self.c
        fi = c.fig_index.get(id(b))
        infos = getattr(c, 'fig_info', None) or []
        info = infos[fi] if fi is not None and fi < len(infos) else {}
        base = self.limit_mm()
        if grid_cols:
            width = (base - GRID_GAP_MM * (grid_cols - 1)) / grid_cols
        else:
            width = min(info.get('display_width_mm') or base * b['width_pct'] / 100, base)
        max_h = 150.0 if self.cur_landscape else 230.0
        if b['kind'] == 'image':
            max_h = min(max_h, (c.landscape_h_mm if self.cur_landscape else c.content_h_mm) * BITMAP_MAX_H_FRACTION)
        return width, max_h

    def caption_p(self, cont, b):
        if not (b.get('caption') or b.get('number')):
            return None
        p = self.new_p(cont, 'DM Figure Caption')
        if b.get('number'): self.run(p._p, b['number'] + '  ', {'b': True, 'color': self.col['primary_dark']})
        self.inline(p._p, b['caption_inline'])
        return p

    def figure(self, cont, b):
        fi = self.c.fig_index.get(id(b))
        eid = self.c.fig_ids[fi] if fi is not None else ''
        p = self.new_p(cont, 'DM Figure')
        w, mh = self.fig_dims(b)
        self.picture(p, b, w, mh)
        cp = self.caption_p(cont, b)
        if eid: self.bookmark(cp or p, eid)

    def grid(self, cont, cols, items):
        if not items:
            return
        self.flush_break_for_table(cont)
        limit = self.limit_mm()
        rows = (len(items) + cols - 1) // cols
        tbl = cont.add_table(rows=rows, cols=cols)
        self.tbl_props(tbl, 'grid', layout='fixed', width_tw=mm_tw(limit), cell_mar=(0, GRID_GAP_MM / 2))
        col_tw = mm_tw(limit / cols)
        self.grid_cols(tbl, [col_tw] * cols)
        w, _ = self.fig_dims(items[0], grid_cols=cols)
        for r0 in range(rows):
            row = tbl.rows[r0]; self.row_flags(row)
            for j in range(cols):
                cell = row.cells[j]; self.set_tcw(cell, col_tw); self.valign(cell, 'top')
                k = r0 * cols + j
                if k >= len(items):
                    continue
                b = items[k]
                fi = self.c.fig_index.get(id(b)); eid = self.c.fig_ids[fi] if fi is not None else ''
                p = self.new_p(cell, 'DM Figure'); p.paragraph_format.space_before = Pt(0)
                _, mh = self.fig_dims(b, grid_cols=cols)
                self.picture(p, b, w - GRID_GAP_MM, mh)
                cp = self.caption_p(cell, b)
                if cp: cp.paragraph_format.space_after = Pt(mm_pt(2))
                if eid: self.bookmark(cp or p, eid)
        self.spacer(cont, mm_pt(4))

    def callout(self, cont, b):
        self.flush_break_for_table(cont)
        tok = self.t['callout'].get(b['callout']) or self.t['callout']['note']
        bg = self.col.get(tok['bg'], tok['bg']); bd = self.col.get(tok['border'], tok['border'])
        limit = self.limit_mm()
        tbl = cont.add_table(rows=1, cols=1)
        self.tbl_props(tbl, 'callout', borders={'left': (36, bd)}, layout='fixed', width_tw=mm_tw(limit), cell_mar=(3, 4))
        self.grid_cols(tbl, [mm_tw(limit)])
        cell = tbl.rows[0].cells[0]
        self.set_tcw(cell, mm_tw(limit)); self.cell_shd(cell, bg); self.cell_borders(cell, {'left': (36, bd)})
        self.row_flags(tbl.rows[0])
        for i, sb in enumerate(b['blocks']):
            if sb['kind'] == 'p':
                p = self.new_p(cell)
                p.paragraph_format.space_after = Pt(0 if i == len(b['blocks']) - 1 else mm_pt(self.sp['para']))
                self.inline(p._p, sb['inline'])
            elif sb['kind'] in ('ul', 'ol'):
                self.lst(cell, sb)
        self.stats['callouts'] += 1
        self.spacer(cont, mm_pt(5))

    def code(self, cont, b):
        lines = b['text'].split('\n')
        keep = len(lines) <= LONG_CODE_LINES
        if b.get('lang'):
            p = self.new_p(cont, CODE_STYLE)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_before = Pt(mm_pt(2)); p.paragraph_format.keep_with_next = True
            self.run(p._p, b['lang'], {'size': self.size['header_footer'] + 1, 'color': self.col['muted'], 'caps': True})
        for i, ln in enumerate(lines):
            p = self.new_p(cont, CODE_STYLE)
            if i == 0 and not b.get('lang'): p.paragraph_format.space_before = Pt(mm_pt(2))
            if i == len(lines) - 1: p.paragraph_format.space_after = Pt(mm_pt(5))
            elif keep: p.paragraph_format.keep_with_next = True
            self.run(p._p, ln.replace('\t', '    ') or ' ')
        self.stats['code_blocks'] += 1

    # ---------- 页眉页脚与分节
    def hf_cell(self, cell, tpl, align):
        p = self.new_p(cell, 'DM Header')
        p.alignment = align
        f = {'color': self.col['subtle'], 'size': self.size['header_footer']}
        for part in re.split(r'(\{logo\}|\{page\}|\{pages\})', fill(tpl or '', self.c.vals)):
            if part == '{logo}':
                if os.path.isfile(self.c.logo_path):
                    run = p.add_run(); run.add_picture(common.logo_for_print(self.c), height=Mm(6))
                    for e in run._r.findall('.//' + qn('wp:docPr')): e.set('name', LOGO_NAME)
                else:
                    self.run(p._p, self.c.vals['brand'], f)
            elif part == '{page}':
                self.field(p._p, 'PAGE', '1', f)
            elif part == '{pages}':
                self.field(p._p, 'NUMPAGES', '1', f)
            elif part:
                self.run(p._p, part, f)

    def hf(self, container, spec, kind, width_mm, rule=False):
        container.is_linked_to_previous = False
        for ch in list(container._element): container._element.remove(ch)
        tbl = container.add_table(rows=1, cols=2, width=Mm(width_mm))
        self.tbl_props(tbl, kind, layout='fixed', width_tw=mm_tw(width_mm), cell_mar=(0, 0))
        left_tw = mm_tw(width_mm / 2); right_tw = mm_tw(width_mm) - left_tw
        self.grid_cols(tbl, [left_tw, right_tw])
        cells = tbl.rows[0].cells
        for j, (tpl, al, tw) in enumerate((((spec or {}).get('left'), WD_ALIGN_PARAGRAPH.LEFT, left_tw), ((spec or {}).get('right'), WD_ALIGN_PARAGRAPH.RIGHT, right_tw))):
            self.set_tcw(cells[j], tw); self.valign(cells[j])
            self.hf_cell(cells[j], tpl, al)
        p = container.add_paragraph(style='DM Header')   # 表后必须有段落；页眉在这一段加主色下边线（§8）
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = Pt(2)
        if rule: self.bottom_rule(p._p.get_or_add_pPr(), self.col['primary'], sz=6, space=1)

    def section_hf(self, sec, landscape, first=False):
        width = self.c.landscape_mm if landscape else self.c.content_mm
        self.hf(sec.header, self.c.profile.get('header'), 'header', width, rule=True)
        self.hf(sec.footer, self.c.profile.get('footer'), 'footer', width)
        sec.different_first_page_header_footer = bool(first)
        if first:   # 封面不印页眉页脚（layout.md §2），页码仍含封面
            for part in (sec.first_page_header, sec.first_page_footer):
                part.is_linked_to_previous = False
                for ch in list(part._element): part._element.remove(ch)
                part.add_paragraph().paragraph_format.space_after = Pt(0)

    def new_section(self, landscape):
        sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        self.page(sec, landscape)
        self.cur_landscape = landscape
        self.section_hf(sec, landscape)
        self.pending_break = False
        self.stats['sections'] += 1
        if landscape: self.stats['landscape_sections'] += 1

    # ---------- 封面与前置页
    def logo_p(self, cont, height_mm, align=WD_ALIGN_PARAGRAPH.LEFT):
        p = self.new_p(cont)
        p.alignment = align; p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
        if os.path.isfile(self.c.logo_path):
            run = p.add_run(); run.add_picture(common.logo_for_print(self.c), height=Mm(height_mm))
            for e in run._r.findall('.//' + qn('wp:docPr')): e.set('name', LOGO_NAME)
        else:
            self.run(p._p, self.c.vals['brand'], {'b': True})
        return p

    def title_lines(self, width_mm):
        from backend_html import balanced_title
        return [html.unescape(x) for x in balanced_title(self.c, self.c.title, width_mm).split('<br/>')]

    def cover(self):
        c = self.c; d = self.d; col = self.col; sz = self.size
        if c.cover == 'marketing':
            self.logo_p(d, 14)
            p = self.new_p(d); p.paragraph_format.space_before = Pt(mm_pt(62)); p.paragraph_format.line_spacing = 1.25; p.paragraph_format.space_after = Pt(0)
            self.run(p._p, '\n'.join(self.title_lines(c.content_mm - 28)), {'b': True, 'size': sz['cover_title']})
            if c.subtitle:
                p = self.new_p(d); p.paragraph_format.space_before = Pt(mm_pt(6)); p.paragraph_format.line_spacing = 1.3
                self.run(p._p, tilde(c, c.subtitle), {'size': sz['cover_subtitle'], 'color': col['primary_dark']})
            from docmark_parse import parse_inline, Diag
            meta = []
            if c.party: meta.append([{'t': 'text', 'v': f'{c.party_label}：{c.party}'}])
            meta += [[{'t': 'text', 'v': f'版本：v{c.version}'}], [{'t': 'text', 'v': f'日期：{c.date}'}]]
            meta += [parse_inline(x, 0, Diag()) for x in c.cover_meta]
            p = self.new_p(d); p.paragraph_format.space_before = Pt(mm_pt(4)); p.paragraph_format.line_spacing = 1.8
            for i, m in enumerate(meta):
                if i: self.br(p._p)
                self.inline(p._p, m, {'size': sz['cover_meta'], 'color': col['ink_2']}, where='cover')
            if c.cover_badge:
                p = self.new_p(d); p.paragraph_format.space_before = Pt(mm_pt(4))
                self.run(p._p, ' ' + fill(c.cover_badge, c.vals) + ' ', {'b': True, 'color': col['white'], 'shd': col['primary']})
            # 底部留白按有无角标自适应：角标块本身约占 10 mm，固定 64 mm 会把封面撑到第二页（LibreOffice 转换时可见，DX8）
            p = self.new_p(d); p.paragraph_format.space_before = Pt(mm_pt(64 - (10 if c.cover_badge else 0)))
            self.run(p._p, c.vals['website'], {'size': sz['cover_meta'], 'color': col['ink_2']})
        else:
            cls = c.profile.get('classification')
            half = mm_tw(c.content_mm / 2)
            top = d.add_table(rows=1, cols=2)
            self.tbl_props(top, 'cover', borders={'top': (48, col['primary'])}, layout='fixed', width_tw=mm_tw(c.content_mm), cell_mar=(2.5, 0))
            self.grid_cols(top, [half, mm_tw(c.content_mm) - half])
            cl, cr = top.rows[0].cells
            for cell, tw in ((cl, half), (cr, mm_tw(c.content_mm) - half)): self.set_tcw(cell, tw); self.valign(cell)
            self.logo_p(cl, 9)
            if cls:
                p = self.new_p(cr); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_after = Pt(0)
                self.run(p._p, f' {cls} ', {'size': sz['caption'], 'color': col['danger'], 'b': True})
            p = self.new_p(d); p.paragraph_format.space_before = Pt(mm_pt(44)); p.paragraph_format.space_after = Pt(mm_pt(4))
            if c.type_name: self.run(p._p, c.type_name, {'b': True, 'size': sz['cover_subtitle'], 'color': col['primary']})
            p = self.new_p(d); p.paragraph_format.line_spacing = 1.25; p.paragraph_format.space_after = Pt(mm_pt(4))
            self.run(p._p, '\n'.join(self.title_lines(c.content_mm)), {'b': True, 'size': sz['cover_title']})
            p = self.spacer(d, mm_pt(6)); p.paragraph_format.right_indent = Mm(c.content_mm - 28)
            self.bottom_rule(p._p.get_or_add_pPr(), col['primary'], sz=24, space=1)
            if c.subtitle:
                self.run(self.new_p(d)._p, tilde(c, c.subtitle), {'size': sz['cover_subtitle'], 'color': col['muted']})
            if c.profile.get('doc_control_table', True):
                rows = []
                if c.meta.get('doc_no'): rows.append(('文档编号', c.meta['doc_no']))
                rows.append(('版本', 'v' + c.version))
                if c.status: rows.append(('状态', STATUS_CN.get(c.status, c.status)))
                if c.owner: rows.append(('负责人', c.owner))
                if c.reviewers: rows.append(('评审人', '\n'.join(reviewer_text(r) for r in c.reviewers)))
                rows.append(('日期', c.date))
                if c.party and c.meta.get('client'): rows.append(('对象', c.party))
                if cls: rows.append(('密级', cls))
                self.spacer(d, mm_pt(36))
                tbl = d.add_table(rows=len(rows) + 1, cols=2)
                self.tbl_props(tbl, 'control', borders={'insideH': (4, col['border']), 'bottom': (4, col['border'])}, layout='fixed', width_tw=mm_tw(c.content_mm))
                ws = [mm_tw(32), mm_tw(c.content_mm) - mm_tw(32)]
                self.grid_cols(tbl, ws)
                hdr = tbl.rows[0]; self.row_flags(hdr, header=True)
                for cell, tw in zip(hdr.cells, ws): self.set_tcw(cell, tw)
                m = hdr.cells[0].merge(hdr.cells[1])
                self.cell_borders(m, {'bottom': (12, col['primary'])})
                self.run(self.new_p(m, 'DM Table Text')._p, '文档控制', {'b': True, 'color': col['muted'], 'size': sz['caption']})
                for i, (k, v) in enumerate(rows):
                    r = tbl.rows[i + 1]; self.row_flags(r)
                    for cell, tw in zip(r.cells, ws): self.set_tcw(cell, tw)
                    self.cell_shd(r.cells[0], col['surface_2'])
                    self.run(self.new_p(r.cells[0], 'DM Table Text')._p, k, {'color': col['muted']})
                    self.run(self.new_p(r.cells[1], 'DM Table Text')._p, v)
            p = self.new_p(d); p.paragraph_format.space_before = Pt(mm_pt(6))
            self.run(p._p, f'{c.vals["brand"]} · {c.vals["website"]}', {'size': sz['caption'], 'color': col['muted']})
        self.pending_break = True

    def front_title(self, text):
        p = self.new_p(self.d, 'DM Front Title')
        self.run(p._p, text)
        return p

    def revision_page(self):
        c = self.c
        if not c.profile.get('revision_page') or not c.revisions:
            return
        self.front_title('修订记录')
        tbl = self.d.add_table(rows=len(c.revisions) + 1, cols=4)
        self.tbl_props(tbl, 'control', borders={s: (4, self.col['border']) for s in BORDER_SIDES}, layout='fixed', width_tw=mm_tw(c.content_mm))
        ws = [mm_tw(c.content_mm * x) for x in (0.12, 0.16, 0.14, 0.58)]
        self.grid_cols(tbl, ws)
        self.row_flags(tbl.rows[0], header=True)
        for i, h in enumerate(('版本', '日期', '作者', '变更摘要')):
            cell = tbl.rows[0].cells[i]; self.set_tcw(cell, ws[i]); self.cell_shd(cell, self.col['surface'])
            self.run(self.new_p(cell, 'DM Table Text')._p, h, {'b': True})
        for ri, r in enumerate(c.revisions):
            row = tbl.rows[ri + 1]; self.row_flags(row)
            v = str(r.get('version') or ''); v = v if v.lower().startswith('v') else 'v' + v
            for i, val in enumerate((v, r.get('date', ''), r.get('author', ''), tilde(c, r.get('summary', '')))):
                self.set_tcw(row.cells[i], ws[i])
                self.run(self.new_p(row.cells[i], 'DM Table Text')._p, str(val))
        self.pending_break = True

    def summary_page(self):
        if self.c.features.get('summary_block') is False:
            return
        blocks = [b for b in self.doc.blocks if b.get('in_summary') and b['kind'] != 'summary']
        if not blocks:
            return
        self.front_title('摘要')
        for b in blocks:
            self.block(self.d, b)
        self.pending_break = True

    def toc(self):
        c = self.c
        if not c.toc_headings:
            return
        self.front_title('目录')
        n = len(c.toc_headings)
        for k, i in enumerate(c.toc_headings):
            h = self.doc.headings[i]
            p = self.new_p(self.d, f'toc {min(h["level"], 2)}')
            if k == 0:
                self.fld_char(p._p, 'begin', dirty=True); self.instr(p._p, f'TOC \\o "1-{max(1, min(c.toc_levels, 9))}" \\h \\z \\u'); self.fld_char(p._p, 'separate')
            bm = self.bm[c.heading_ids[i]]
            hl = OxmlElement('w:hyperlink'); hl.set(qn('w:anchor'), bm); hl.set(qn('w:history'), '1')
            lab = heading_label(h)
            if lab: self.run(hl, lab + '  ')
            self.inline(hl, h['inline'], {}, where='toc')
            p._p.append(hl)
            self.run(p._p, None).append(OxmlElement('w:tab'))
            self.field(p._p, f'PAGEREF {bm} \\h', '')
            if k == n - 1: self.fld_char(p._p, 'end')
            self.toc_entries += 1
        self.stats['toc'] = True
        self.pending_break = True

    # ---------- 正文
    def body(self):
        c = self.c; doc = self.doc
        first_h1 = True
        grid = None
        for b in doc.blocks:
            if b.get('in_summary') or b['kind'] == 'summary':
                continue
            k = b['kind']
            if k == 'heading':
                h = doc.headings[b['heading']]
                if h['level'] == 1:
                    if c.h1_mode == 'always' and not first_h1: self.pending_break = True
                    first_h1 = False
                p = self.new_p(self.d, f'Heading {min(h["level"], 4)}')
                lab = heading_label(h)
                if lab: self.run(p._p, lab + '  ')
                self.inline(p._p, h['inline'], {})
                self.bookmark(p, c.heading_ids[b['heading']])
                self.stats['headings'] += 1
                continue
            if k == 'landscape':
                if c.features.get('landscape') is False:
                    continue
                if b['edge'] == 'start' and not self.cur_landscape: self.new_section(True)
                elif b['edge'] == 'end' and self.cur_landscape: self.new_section(False)
                continue
            if k == 'grid':   # 与 HTML、飞书后端一致：grid 只收图，结束标记处整体输出
                if b['edge'] == 'start':
                    grid = (b['cols'], [])
                elif grid is not None:
                    self.grid(self.d, grid[0], grid[1]); grid = None
                continue
            if grid is not None and k in ('figure', 'image') and b.get('grid'):
                grid[1].append(b); continue
            self.block(self.d, b)
        if grid is not None:
            self.grid(self.d, grid[0], grid[1])
        if c.profile.get('disclaimer'):
            p = self.new_p(self.d); p.paragraph_format.space_before = Pt(mm_pt(10))
            self.run(p._p, fill(c.profile['disclaimer'], c.vals), {'size': self.size['footnote'], 'color': self.col['subtle']})

    # ---------- 收尾：脚注部件、设置、属性
    def finish(self):
        c = self.c
        if self.fn_xml:
            from lxml import etree
            sep = '<w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:{}/></w:r></w:p>'
            root = parse_xml(f'<w:footnotes {nsdecls("w", "r")}><w:footnote w:type="separator" w:id="-1">{sep.format("separator")}</w:footnote>'
                             f'<w:footnote w:type="continuationSeparator" w:id="0">{sep.format("continuationSeparator")}</w:footnote></w:footnotes>')
            for wid, p in self.fn_xml:
                fn = OxmlElement('w:footnote'); fn.set(qn('w:id'), str(wid)); fn.append(p); root.append(fn)
            blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
            self.d.part.relate_to(Part(PackURI('/word/footnotes.xml'), CT_FOOTNOTES, blob, self.d.part.package), RT.FOOTNOTES)
            fpr = set_child(self.d.settings.element, 'w:footnotePr', SEQ_SETTINGS)
            for i in ('-1', '0'):
                e = OxmlElement('w:footnote'); e.set(qn('w:id'), i); fpr.append(e)
        if self.stats['toc']:
            set_child(self.d.settings.element, 'w:updateFields', SEQ_SETTINGS, {'w:val': 'true'})
        cp = self.d.core_properties
        md = common.doc_metadata(c)   # 与 PDF 元数据同一口径（layout.md §11）
        cp.title = md['Title']; cp.author = md['Author']; cp.subject = md['Subject'] or ''; cp.keywords = md['Keywords']
        cp.category = c.type or ''
        cp.comments = f'{common.RENDERER_NAME} {common.RENDERER_VERSION}'
        cp.last_modified_by = common.RENDERER_NAME
        cp.language = c.meta.get('language') or 'zh-CN'
        cp.revision = 1
        now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
        cp.created = now; cp.modified = now

    def build(self, out_path):
        self.setup()
        self.section_hf(self.d.sections[0], False, first=True)
        self.cover()
        self.revision_page()
        self.summary_page()
        self.toc()
        self.body()
        self.finish()
        self.d.save(out_path)
        self.stats['bytes'] = os.path.getsize(out_path)
        self.stats['_grids'] = self.grids
        self.stats['_toc_entries'] = self.toc_entries
        self.stats['_fn_ids'] = sorted(self.fn_seq)
        self.stats['_bm'] = dict(self.bm)
        return self.stats


def build_docx(ctx, out_path):
    """写 out_path，返回统计（以下划线开头的键只给 readback 用，不进 render.json）。"""
    os.makedirs(os.path.dirname(out_path) or '.', exist_ok=True)
    return DocxBuilder(ctx).build(out_path)


# ---------- 读回校验（layout.md §13）
def _texts(el, skip_super=False):
    out = []
    for r in el.iter(qn('w:r')):
        if skip_super:
            rpr = r.find(qn('w:rPr'))
            if rpr is not None and (rpr.find(qn('w:vertAlign')) is not None or rpr.find(qn('w:rStyle')) is not None):
                continue
        out.extend(t.text or '' for t in r.iter(qn('w:t')))
    return ''.join(out)


def _visible_borders(tbl):
    bad = []
    for e in tbl.iter(qn('w:tblBorders'), qn('w:tcBorders')):
        for s in e:
            if (s.get(qn('w:val')) or 'nil') not in ('nil', 'none'):
                bad.append(s.tag.split('}')[1])
    return bad


def readback(ctx, path, stats):
    """ZIP + lxml 核对精确 XML，python-docx 核对对象模型。返回 (readback, layout_issues)。"""
    checks, issues = [], []

    def chk(rule, name, ok, expected, actual, severity='必改', message=None):
        checks.append({'rule': rule, 'name': name, 'ok': bool(ok), 'severity': severity, 'expected': expected, 'actual': actual})
        if not ok:
            issues.append({'rule': rule, 'severity': severity, 'page': None, 'target': 'docx',
                           'message': message or f'docx 读回：{name}（期望 {expected}，实际 {actual}）'})

    try:
        d = Document(path)
        zf = zipfile.ZipFile(path)
        names = set(zf.namelist())
        doc_xml = zf.read('word/document.xml').decode('utf-8')
        settings_xml = zf.read('word/settings.xml').decode('utf-8') if 'word/settings.xml' in names else ''
        fn_xml = zf.read('word/footnotes.xml').decode('utf-8') if 'word/footnotes.xml' in names else ''
        ct_xml = zf.read('[Content_Types].xml').decode('utf-8')
        rels_xml = zf.read('word/_rels/document.xml.rels').decode('utf-8') if 'word/_rels/document.xml.rels' in names else ''
        zf.close()
    except Exception as ex:  # noqa: BLE001
        chk('DX9', '重新打开 docx', False, '可打开', str(ex)[:160])
        return {'ok': False, 'all_passed': False, 'checks': checks, 'highlights': []}, issues
    doc = ctx.doc
    body = d.element.body
    sid = {s.name: s.style_id for s in d.styles}
    kind_of = {sid.get(v): k for k, v in TABLE_STYLES.items() if sid.get(v)}

    def pstyle(p):
        e = p.find(qn('w:pPr') + '/' + qn('w:pStyle'))
        return e.get(qn('w:val')) if e is not None else None

    def tkind(t):
        e = t.find(qn('w:tblPr') + '/' + qn('w:tblStyle'))
        return kind_of.get(e.get(qn('w:val'))) if e is not None else None

    # DX1 标题
    level_of = {sid[f'Heading {i}']: i for i in range(1, 10) if f'Heading {i}' in sid}
    got = [(level_of[pstyle(p)], re.sub(r'\s+', '', _texts(p, skip_super=True))) for p in body.iter(qn('w:p')) if pstyle(p) in level_of]
    exp = [(min(h['level'], 4), re.sub(r'\s+', '', tilde(ctx, (heading_label(h) or '') + common.dp.inline_plain(h['inline'])))) for h in doc.headings]
    chk('DX1', '标题数与层级（Word 内置 Heading 样式）', [g[0] for g in got] == [e[0] for e in exp], f'{len(exp)} 个', f'{len(got)} 个')
    bad = [(e[1], g[1]) for e, g in zip(exp, got) if e[1] != g[1]]
    chk('DX1', '标题文字（含编号）与解析器一致', not bad and len(exp) == len(got), '一致', str(bad[:2]) if bad else '一致')

    # DX2 / DX3 内容表
    tbls = list(body.iter(qn('w:tbl')))
    content = [t for t in tbls if tkind(t) == 'table']
    exp_tables = sum(1 for b in doc.tables if not b.get('error'))
    chk('DX2', '内容表数', len(content) == exp_tables == stats.get('tables'), exp_tables, len(content))
    unknown = [i + 1 for i, t in enumerate(tbls) if tkind(t) is None]
    chk('DX2', '表格样式标记（DM Table 等）', not unknown, '全部有', f'缺 {unknown[:8]}' if unknown else '全部有')
    if ctx.features.get('thead_repeat') is not False:
        no_hdr = [i + 1 for i, t in enumerate(content) if t.find(qn('w:tr') + '/' + qn('w:trPr') + '/' + qn('w:tblHeader')) is None]
        chk('DX2', '内容表首行「标题行重复」', not no_hdr, '全部有 tblHeader', f'缺 {no_hdr[:8]}' if no_hdr else '全部有')
    split = [i + 1 for i, t in enumerate(content) if any(tr.find(qn('w:trPr') + '/' + qn('w:cantSplit')) is None for tr in t.findall(qn('w:tr')))]
    chk('DX2', '内容表每行不跨页断开（cantSplit）', not split, '全部行', f'表 {split[:8]}' if split else '全部行')
    bad_fit = []
    for i, t in enumerate(content):
        tp = t.find(qn('w:tblPr')); lay = tp.find(qn('w:tblLayout')); tw = tp.find(qn('w:tblW'))
        if lay is None or lay.get(qn('w:type')) != 'autofit' or tw is None or tw.get(qn('w:type')) != 'auto' or t.find('.//' + qn('w:tcW')) is not None:
            bad_fit.append(i + 1)
    chk('DX3', '内容表列宽自适应（tblLayout=autofit、tblW=auto、无 tcW）', not bad_fit, '全部满足', f'不满足 {bad_fit[:8]}' if bad_fit else '全部满足')
    bad_grid = []
    for i, (t, (limit_tw, want)) in enumerate(zip(content, stats.get('_grids') or [])):
        cols = [int(g.get(qn('w:w'))) for g in t.find(qn('w:tblGrid')).findall(qn('w:gridCol'))]
        tot = sum(cols) or 1
        if abs(tot - limit_tw) > limit_tw * 0.02 or len(cols) != len(want) or any(abs(a / tot - b / (sum(want) or 1)) > 0.01 for a, b in zip(cols, want)):
            bad_grid.append(i + 1)
    chk('DX3', '内容表初始网格 = common.table_widths × 当前版心宽', not bad_grid, '全部一致', f'不一致 {bad_grid[:8]}' if bad_grid else '全部一致')

    # DX4 图
    pics = [e for e in body.iter(qn('wp:docPr')) if e.get('name') != LOGO_NAME]
    chk('DX4', '正文图片数', len(pics) == stats.get('images'), stats.get('images'), len(pics))
    missing = stats.get('missing_images') or []
    chk('DX4', '图全部嵌入（doc-figures PNG / 位图）', not missing and stats.get('images') == len(doc.figures), len(doc.figures), f'嵌入 {stats.get("images")}，缺 {len(missing)}',
        message=('docx 有 %d 张图未嵌入：%s（先运行 doc-figures build 生成 PNG）' % (len(missing), '、'.join(m['expected'] for m in missing[:5]))) if missing else None)

    # DX5 页眉页脚
    hf_bad = []
    for si, sec in enumerate(d.sections):
        for label, part, kind in (('页眉', sec.header, 'header'), ('页脚', sec.footer, 'footer')):
            if part.is_linked_to_previous:
                if si == 0: hf_bad.append(f'第 1 节{label}缺失')
                continue
            ts = [t for t in part._element.iter(qn('w:tbl')) if tkind(t) == kind]
            if not ts:
                hf_bad.append(f'第 {si + 1} 节{label}不是两列表格')
            for t in ts:
                if len(t.findall(qn('w:tr') + '/' + qn('w:tc'))) != 2: hf_bad.append(f'第 {si + 1} 节{label}不是 2 列')
                vb = _visible_borders(t)
                if vb: hf_bad.append(f'第 {si + 1} 节{label}有边框 {sorted(set(vb))}')
                lay = t.find(qn('w:tblPr') + '/' + qn('w:tblLayout'))
                if lay is None or lay.get(qn('w:type')) != 'fixed': hf_bad.append(f'第 {si + 1} 节{label}不是 fixed 布局')
                limit = int(round((sec.page_width - sec.left_margin - sec.right_margin) / 635))   # EMU → twips
                tw = t.find(qn('w:tblPr') + '/' + qn('w:tblW'))
                tw_v = int(tw.get(qn('w:w'))) if tw is not None and tw.get(qn('w:type')) == 'dxa' else None
                grid_w = sum(int(g.get(qn('w:w'))) for g in t.find(qn('w:tblGrid')).findall(qn('w:gridCol')))
                tcw = sum(int(e.get(qn('w:w'))) for e in t.findall(qn('w:tr') + '/' + qn('w:tc') + '/' + qn('w:tcPr') + '/' + qn('w:tcW')))
                if tw_v is None or abs(tw_v - limit) > 3 or abs(grid_w - limit) > 3 or abs(tcw - limit) > 3:
                    hf_bad.append(f'第 {si + 1} 节{label}表宽 tblW={tw_v} tblGrid={grid_w} tcW={tcw} ≠ 版心 {limit} twips')
        foot_xml = ''.join(_instr(p._element) for p in [sec.footer] if not sec.footer.is_linked_to_previous)
        spec = ''.join((ctx.profile.get('footer') or {}).values())
        if not sec.footer.is_linked_to_previous:
            if '{page}' in spec and ' PAGE ' not in foot_xml: hf_bad.append(f'第 {si + 1} 节页脚缺 PAGE 域')
            if '{pages}' in spec and ' NUMPAGES ' not in foot_xml: hf_bad.append(f'第 {si + 1} 节页脚缺 NUMPAGES 域')
    s0 = d.sections[0]
    if not s0.different_first_page_header_footer:
        hf_bad.append('第 1 节未设首页不同（封面会印页眉页脚）')
    else:
        for label, part in (('页眉', s0.first_page_header), ('页脚', s0.first_page_footer)):
            el = part._element
            if _texts(el).strip() or el.find('.//' + qn('w:tbl')) is not None or el.find('.//' + qn('w:drawing')) is not None:
                hf_bad.append(f'封面首页{label}不为空')
    chk('DX5', '页眉页脚：两列无边框 fixed 表、表宽等于各节版心、页码域、封面首页为空', not hf_bad, '全部满足', '；'.join(hf_bad[:6]) or '全部满足')

    # DX6 语法残留（所有 w:t，跳过代码块段落）
    leaks = []
    roots = [('正文', body)]
    for sec in d.sections:
        for part in (sec.header, sec.footer):
            if not part.is_linked_to_previous: roots.append(('页眉页脚', part._element))
    if fn_xml: roots.append(('脚注', parse_xml(fn_xml.encode('utf-8'))))
    code_id = sid.get(CODE_STYLE)
    for label, root in roots:
        for p in root.iter(qn('w:p')):
            if code_id and pstyle(p) == code_id:
                continue
            tx = ''.join(t.text or '' for t in p.iter(qn('w:t')))
            m = LEAK_RE.search(tx)
            if m: leaks.append(f'{label}：…{tx[max(0, m.start() - 12):m.end() + 12]}…')
    chk('DX6', '无 ⟪ ⟫ 与 DocMark 语法残留', not leaks, '0 处', '；'.join(leaks[:4]) or '0 处')

    # DX7 核心属性
    cp = d.core_properties
    want = {k.lower(): (v or '') for k, v in common.doc_metadata(ctx).items() if k in ('Title', 'Author', 'Subject', 'Keywords')}
    got = {'title': cp.title or '', 'author': cp.author or '', 'subject': cp.subject or '', 'keywords': cp.keywords or ''}
    diff7 = [k for k in want if got[k] != want[k] or not got[k].strip()]
    chk('DX7', '核心属性 Title / Author / Subject / Keywords 与 PDF 元数据口径逐项相同且非空', not diff7, want, {k: got[k] for k in diff7} if diff7 else '一致')

    # DX8 目录、书签链接、脚注、高亮、横向节（建议级）
    want_toc = bool(ctx.toc_headings)
    has_toc = ' TOC \\o' in doc_xml
    upd = re.search(r'<w:updateFields w:val="(true|1)"', settings_xml) is not None
    chk('DX8', '目录域、条目数与打开时更新提示', has_toc == want_toc and (not want_toc or (upd and stats.get('_toc_entries') == len(ctx.toc_headings))),
        f'TOC={want_toc} 条目 {len(ctx.toc_headings)}', f'TOC={has_toc} 条目 {stats.get("_toc_entries")} updateFields={upd}', severity='建议')
    marks = {e.get(qn('w:name')) for e in body.iter(qn('w:bookmarkStart'))}
    dangling = sorted({e.get(qn('w:anchor')) for e in body.iter(qn('w:hyperlink')) if e.get(qn('w:anchor')) and e.get(qn('w:anchor')) not in marks})
    chk('DX8', '内部链接（交叉引用、目录）都有书签目标', not dangling, '0 个悬空', str(dangling[:5]) if dangling else '0 个悬空', severity='建议')
    bm_map = stats.get('_bm') or {}
    visible = {n for n in marks if n and not n.startswith('_')}
    want_vis = {bm_map.get(ctx.heading_ids[i]) for i in ctx.bookmark_headings}
    chk('DX1', f'可见书签 = 一至{ctx.bookmark_levels}级标题（{len(ctx.bookmark_headings)} 个），其余目标为隐藏书签', visible == want_vis and len(visible) == len(ctx.bookmark_headings),
        len(ctx.bookmark_headings), f'{len(visible)} 个，多 {sorted(visible - want_vis)[:3]}，少 {sorted(x for x in want_vis - visible if x)[:3]}')
    all_marks = marks | set(re.findall(r'<w:bookmarkStart [^>]*w:name="([^"]+)"', fn_xml))
    all_links = {e.get(qn('w:anchor')) for e in body.iter(qn('w:hyperlink'))} | set(re.findall(r'<w:hyperlink [^>]*w:anchor="([^"]+)"', fn_xml))
    resolved = [r for r in doc.refs if r.get('resolved')]
    bad_refs = [r['target'] for r in resolved if not (bm_map.get(ctx.anchor_elem.get(r['target'])) in all_marks and bm_map.get(ctx.anchor_elem.get(r['target'])) in all_links)]
    chk('DX8', f'交叉引用：已解析引用 {len(resolved)} 条逐条有内部链接与书签目标（含实体锚点）', not bad_refs, f'{len(resolved)} 条', f'缺 {sorted(set(bad_refs))[:5]}' if bad_refs else '全部有', severity='建议')
    refs = len(re.findall(r'<w:footnoteReference ', doc_xml))
    bodies = len(re.findall(r'<w:footnote w:id="[1-9]\d*"', fn_xml))
    fn_refs_marks = len(re.findall(r'<w:footnoteRef/>', fn_xml))
    noterefs = doc_xml.count(' NOTEREF _Ref_dmfn')
    exp_ids = sorted(ctx.fn_num)
    if refs or fn_xml:
        ct_ok = 'PartName="/word/footnotes.xml"' in ct_xml and CT_FOOTNOTES in ct_xml
        rel_ok = re.search(r'Type="[^"]*/footnotes"[^>]*Target="footnotes.xml"|Target="footnotes.xml"[^>]*Type="[^"]*/footnotes"', rels_xml) is not None
        note_targets = set(re.findall(r' NOTEREF (_Ref_dmfn\d+) ', doc_xml))
        note_ok = note_targets <= {m for m in marks if m.startswith('_Ref_dmfn')}
        ids_ok = set(re.findall(r'<w:footnoteReference w:id="(\d+)"', doc_xml)) == set(re.findall(r'<w:footnote w:id="([1-9]\d*)"', fn_xml))
        chk('DX8', '脚注部件：Content-Type、文档关系、引用 id 与脚注正文 id 一一对应、NOTEREF 书签存在', ct_ok and rel_ok and note_ok and ids_ok,
            '全部满足', f'content-type={ct_ok} rel={rel_ok} noteref={note_ok} ids={ids_ok}', severity='建议')
    chk('DX8', '原生脚注：引用 = 脚注正文 = 被引用的脚注数，再次引用为 NOTEREF', refs == bodies == fn_refs_marks == stats.get('footnotes') == len(exp_ids)
        and noterefs == stats.get('footnote_repeats') and stats.get('_fn_ids') == exp_ids,
        f'{len(exp_ids)} 个脚注', f'引用 {refs}、正文 {bodies}、NOTEREF {noterefs}', severity='建议')
    # 五种高亮 kind 的底色都算重点高亮（kind 不影响计数口径，只影响底色）
    hl_fills = {hexc(hl_fill(ctx.tokens, k)) for k in ('neutral', 'risk', 'tip', 'warn', 'decision')}
    groups = []
    for p in body.iter(qn('w:p')):
        cur = None
        for r in p.iter(qn('w:r')):
            shd = r.find(qn('w:rPr') + '/' + qn('w:shd'))
            if shd is not None and (shd.get(qn('w:fill')) or '').upper() in hl_fills:
                if cur is None: cur = []; groups.append(cur)
                cur.append(''.join(t.text or '' for t in r.iter(qn('w:t'))))
            elif r.find(qn('w:t')) is not None:
                cur = None
    hl = [''.join(g) for g in groups]
    chk('DX8', '重点高亮组数（tokens callout 底色 + 加粗，5 种 kind 合计）', len(hl) == stats.get('highlights'), stats.get('highlights'), len(hl), severity='建议')
    land = [s for s in d.sections if s.orientation == WD_ORIENT.LANDSCAPE]
    geo = [i + 1 for i, s in enumerate(d.sections) if (s.orientation == WD_ORIENT.LANDSCAPE) != (s.page_width > s.page_height)]
    chk('DX8', '横向节数与纸张方向', len(land) == stats.get('landscape_sections', 0) and not geo, stats.get('landscape_sections', 0), f'{len(land)}，方向不符 {geo}' if geo else len(land), severity='建议')
    return {'ok': all(x['ok'] for x in checks if x['severity'] == '必改'), 'all_passed': all(x['ok'] for x in checks), 'checks': checks, 'highlights': hl}, issues


def _instr(el):
    return ''.join(' ' + (t.text or '').strip() + ' ' for t in el.iter(qn('w:instrText')))


def soffice_check(ctx, path, stats, work_dir, timeout=240):
    """LibreOffice 真实客户端核验：docx → PDF 转换成功、页数、横向页、标题文字可见。
    返回 (soffice 记录或 None（未安装）, layout_issues)。"""
    exe = soffice_bin()
    if not exe or os.environ.get('DOC_RENDER_DOCX_SOFFICE') == '0':
        return None, []
    os.makedirs(work_dir, exist_ok=True)
    profile = tempfile.mkdtemp(prefix='dm-soffice-')
    rec = {'ok': False}
    try:
        try:
            v = subprocess.run([exe, '--version'], capture_output=True, text=True, timeout=60).stdout.strip()
            if v: rec['version'] = v.split('\n')[0][:80]
        except (OSError, subprocess.TimeoutExpired):
            pass
        env = dict(os.environ)
        if not env.get('FONTCONFIG_FILE'): env['FONTCONFIG_FILE'] = fontconfig_file(exe)
        r = subprocess.run([exe, f'-env:UserInstallation=file://{profile}', '--headless', '--convert-to', 'pdf', '--outdir', work_dir, path],
                           capture_output=True, text=True, timeout=timeout, env=env)
        pdf = os.path.join(work_dir, os.path.splitext(os.path.basename(path))[0] + '.pdf')
        if r.returncode != 0 or not os.path.isfile(pdf):
            rec['error'] = f'转换失败（退出码 {r.returncode}）：{(r.stdout + r.stderr)[-200:]}'
        else:
            from pypdf import PdfReader
            rd = PdfReader(pdf)
            rec['pages'] = len(rd.pages)
            rec['landscape_pages'] = sum(1 for pg in rd.pages if float(pg.mediabox.width) > float(pg.mediabox.height))
            if shutil.which('pdftotext'):   # pypdf 取不出 LibreOffice 子集化中文字体的文字，优先 poppler
                raw = subprocess.run(['pdftotext', pdf, '-'], capture_output=True, text=True, timeout=120).stdout
            else:
                raw = '\f'.join((pg.extract_text() or '') for pg in rd.pages)
            text = re.sub(r'\s+', '', raw)
            pages_text = raw.split('')[:len(rd.pages)]
            rec['blank_pages'] = [pi + 1 for pi, pg_text in enumerate(pages_text) if pi > 0 and not page_residue(ctx, pg_text)]
            probes = [re.sub(r'\s+', '', ctx.title)[:10]] + [re.sub(r'\s+', '', ctx.doc.headings[i]['plain'])[:10] for i in range(min(3, len(ctx.doc.headings)))]
            lost = [x for x in probes if x and x not in text]
            errs = []
            # 文字层在不代表字形印出来了：没有嵌入中文字体时页面上中文不可见（本机实测过）。按转换后 PDF 的实际文字判断
            fonts_out = subprocess.run(['pdffonts', pdf], capture_output=True, text=True, timeout=60).stdout if shutil.which('pdffonts') else None
            ferrs, embedded = font_problems(text, fonts_out)
            if embedded: rec['fonts'] = sorted(set(embedded))[:12]
            errs += ferrs
            if rec['pages'] < 2: errs.append(f'页数 {rec["pages"]}')
            if bool(rec['landscape_pages']) != bool(stats.get('landscape_sections')): errs.append(f'横向页 {rec["landscape_pages"]} 与横向节 {stats.get("landscape_sections")} 不符')
            if lost: errs.append(f'转换后找不到文字 {lost}')
            rec['ok'] = not errs
            if errs: rec['error'] = '；'.join(errs)
    except (OSError, subprocess.TimeoutExpired) as ex:
        rec['error'] = f'LibreOffice 无法运行：{ex}'[:200]
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    issues = [] if rec['ok'] else [{'rule': 'DX9', 'severity': '必改', 'page': None, 'target': 'docx', 'message': f'docx 在 LibreOffice 中转换核验失败：{rec.get("error")}'}]
    if rec.get('blank_pages'):
        issues.append({'rule': 'DX8', 'severity': '建议', 'page': rec['blank_pages'][0], 'target': 'docx',
                       'message': f'docx 经 LibreOffice 转换后第 {"、".join(map(str, rec["blank_pages"]))} 页为空白页（只有页眉页脚）；Word 排版可能不同'})
    return rec, issues
