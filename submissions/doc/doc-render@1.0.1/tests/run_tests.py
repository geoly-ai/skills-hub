#!/usr/bin/env python3
"""doc-render 自测。离线：不发布飞书，不改 doc-shared、presales-* 与样张目录（样张复制到系统临时目录再渲染）。
用法：run_tests.py [--quick] [--keep] [--previews <目录>]
  --quick      只跑解析器与不打印 PDF 的用例
  --keep       保留临时目录
  --previews   把三份样张 PDF 的首页与目录页导出 PNG 到该目录
  --docx-previews <目录>  把三份样张与 golden smoke-site 的 docx 经 LibreOffice 转出的 PDF 逐页导出 PNG 到该目录
  --allow-missing-deps  缺 poppler 或 lark-cli 时记 SKIP 而不是 FAIL（默认缺依赖即失败，避免「ALL PASS」掩盖没跑的检查）"""
import glob, json, os, re, shutil, subprocess, sys, tempfile

H = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(H)
SKILLS = os.path.dirname(ROOT)
DOC_SHARED = os.path.join(SKILLS, 'doc-shared')
VENV_PY = os.path.join(ROOT, '.venv', 'bin', 'python')
RENDER = os.path.join(ROOT, 'scripts', 'render.py')
STAGING = os.path.expanduser('~/workspace/docs/doc-skill-platform-2026-09-15/staging/types')
GOLDEN = os.path.expanduser('~/workspace/docs/doc-skill-platform-2026-09-15/golden')
FIXTURE_PACKS = os.path.join(H, 'fixtures', 'packs')
SAMPLE_TYPES = ('prd', 'tech-spec', 'test-cases')

try:
    import pypdf  # noqa: F401
except ImportError:
    if os.path.exists(VENV_PY) and not os.environ.get('DOC_RENDER_TEST_REEXEC'):
        os.environ['DOC_RENDER_TEST_REEXEC'] = '1'
        os.execv(VENV_PY, [VENV_PY, os.path.abspath(__file__)] + sys.argv[1:])
from pypdf import PdfReader  # noqa: E402
import logging; logging.getLogger('pypdf').setLevel(logging.ERROR)  # noqa: E402

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.join(DOC_SHARED, 'scripts')); sys.path.insert(0, os.path.join(DOC_SHARED, 'tests'))
import validate  # noqa: E402
import test_docmark_parse  # noqa: E402

args = sys.argv[1:]
QUICK = '--quick' in args
KEEP = '--keep' in args
PREVIEWS = args[args.index('--previews') + 1] if '--previews' in args else None
DOCX_PREVIEWS = args[args.index('--docx-previews') + 1] if '--docx-previews' in args else None
ALLOW_MISSING = '--allow-missing-deps' in args
fails, skips = [], []
BASE = tempfile.mkdtemp(prefix='doc-render-tests-')


def check(name, cond, detail=''):
    print(('PASS ' if cond else 'FAIL ') + name + (f'  [{str(detail)[:600]}]' if detail and not cond else ''))
    if not cond: fails.append(name)


def skip(name, why):
    print(f'SKIP {name}（{why}）'); skips.append(name)


def missing_dep(name, why):
    if ALLOW_MISSING: skip(name, why)
    else: check(f'{name}（依赖：{why}）', False, '缺依赖；确实无法安装时加 --allow-missing-deps')


def run(cmd, cwd=None, env=None, timeout=600):
    r = subprocess.run(cmd, cwd=cwd, capture_output=True, text=True, env=env, timeout=timeout)
    return r.returncode, r.stdout, r.stderr


def jout(o):
    try:
        i = o.find('{'); return json.loads(o[i:])
    except Exception:
        return {}


def norm(s):
    return re.sub(r'\s+', '', s)


def svg(title, w=1200, px=17, n=5):
    rows = []
    for i in range(n):
        y = 90 + i * 70
        rows.append(f'<rect x="60" y="{y}" width="480" height="48" rx="10" fill="#EEEBFF" stroke="#6C5CE7" stroke-width="1.8"/>'
                    f'<text x="300" y="{y + 30}" font-size="19" text-anchor="middle" fill="#3C2A9E">{title} 节点 {i + 1}</text>'
                    f'<text x="700" y="{y + 30}" class="note" fill="#5F6368">说明 {i + 1}</text>')
    h = 90 + n * 70 + 20
    return (f'<?xml version="1.0"?><svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {w} {h}" width="{w}" height="{h}"><style>.note{{font-size:{px}px}}</style>'
            f'<defs><marker id="arr" markerWidth="11" markerHeight="11" refX="10" refY="5" orient="auto"><polygon points="0,0 11,5 0,10"/></marker></defs>'
            f'<line x1="540" y1="100" x2="660" y2="100" stroke="#5F6368" marker-end="url(#arr)"/>{"".join(rows)}</svg>')


def png_bytes(w, h, rgb=(238, 235, 255)):
    import struct, zlib
    raw = b''.join(b'\x00' + bytes(rgb) * w for _ in range(h))
    def chunk(t, d): return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')


def build_png(svg_path):
    """模拟 doc-figures build 的 2 倍像素 PNG：有 rsvg-convert 时真栅格化（预览好看），否则写同宽高比的纯色 PNG（不引入测试依赖）。"""
    out = svg_path[:-4] + '.png'
    if os.path.exists(out): return
    if shutil.which('rsvg-convert') and subprocess.run(['rsvg-convert', '-w', '2400', svg_path, '-o', out], capture_output=True).returncode == 0:
        return
    m = re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', open(svg_path, encoding='utf-8').read())
    w, h = (float(m.group(1)), float(m.group(2))) if m else (1200.0, 460.0)
    open(out, 'wb').write(png_bytes(600, max(1, int(600 * h / w))))


def build_svgs(run_dir):
    for src in glob.glob(os.path.join(run_dir, 'figures', '*')):
        if src.endswith(('.svg', '.png')): continue
        base = src[:-len('.fig.json')] if src.endswith('.fig.json') else src.rsplit('.', 1)[0]
        if not os.path.exists(base + '.svg'):
            open(base + '.svg', 'w', encoding='utf-8').write(svg(os.path.basename(base)))
        build_png(base + '.svg')


def docx_checks(t, rd, rj, doc, soffice_required=True, expect_all_passed=True, meta_expect=None):
    """独立于 backend_docx.readback 的核对（直接用 python-docx 与 zip 读 XML）。
    meta_expect：核心属性期望值；缺省取 render.json 里 PDF 元数据（同一口径的另一条产出路径）。"""
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn
    dj = rj.get('docx') or {}
    rb = dj.get('readback') or {}
    fails_ = [c_ for c_ in rb.get('checks', []) if not c_['ok']]
    check(f'{t}：docx 已输出，读回无必改级失败（ok）' + ('且全部检查通过（all_passed）' if expect_all_passed else ''),
          bool(dj) and rb.get('ok') is True and (rb.get('all_passed') is True or not expect_all_passed) and rb.get('ok') == all(c_['ok'] for c_ in rb.get('checks', []) if c_.get('severity') == '必改'), fails_[:3] or dj)
    path = os.path.join(rd, dj.get('path', '-'))
    if not os.path.isfile(path):
        return None
    d = DocxDocument(path); body = d.element.body
    sid = {st.name: st.style_id for st in d.styles}
    def pst(p): e = p.find(qn('w:pPr') + '/' + qn('w:pStyle')); return e.get(qn('w:val')) if e is not None else None
    def tst(x): e = x.find(qn('w:tblPr') + '/' + qn('w:tblStyle')); return e.get(qn('w:val')) if e is not None else None
    heads = [p for p in body.iter(qn('w:p')) if pst(p) in {sid.get(f'Heading {i}') for i in range(1, 5)}]
    tbls = [x for x in body.iter(qn('w:tbl')) if tst(x) == sid.get('DM Table')]
    hdr_ok = all(x.find(qn('w:tr') + '/' + qn('w:trPr') + '/' + qn('w:tblHeader')) is not None for x in tbls)
    pics = [e for e in body.iter(qn('wp:docPr')) if (e.get('name') or '').startswith('dm-figure')]
    check(f'{t}：docx 独立核对 Heading 段落 {len(heads)} = 标题 {len(doc.headings)}，内容表 {len(tbls)} = {len(doc.tables)} 且首行标题行重复，图片 {len(pics)} = {len(doc.figures)}',
          len(heads) == len(doc.headings) and len(tbls) == len(doc.tables) and hdr_ok and len(pics) == len(doc.figures))
    hts = [x for x in d.sections[0].header._element.iter(qn('w:tbl'))]
    vis = [b_.get(qn('w:val')) for x in hts for e in x.iter(qn('w:tblBorders'), qn('w:tcBorders')) for b_ in e if b_.get(qn('w:val')) not in ('nil', 'none')]
    check(f'{t}：docx 页眉为 1 行 2 列无边框表格，页脚有 PAGE 域', len(hts) == 1 and len(hts[0].findall(qn('w:tr') + '/' + qn('w:tc'))) == 2 and not vis
          and ' PAGE ' in ''.join(' ' + (x.text or '').strip() + ' ' for x in d.sections[0].footer._element.iter(qn('w:instrText'))), (len(hts), vis))
    cp = d.core_properties
    me = meta_expect or {k: (rj.get('pdf') or {}).get('metadata', {}).get(k) for k in ('Title', 'Author', 'Subject', 'Keywords')}
    got_me = {'Title': cp.title, 'Author': cp.author, 'Subject': cp.subject, 'Keywords': cp.keywords}
    check(f'{t}：docx 核心属性 Title / Author / Subject / Keywords 逐项等于期望（{"PDF 元数据" if not meta_expect else "doc.json 推得"}）', all(me.values()) and got_me == me, (got_me, me))
    hf_bad, land_cov = [], 0
    for si, sec in enumerate(d.sections):
        limit = int(round((sec.page_width - sec.left_margin - sec.right_margin) / 635))
        for part in (sec.header, sec.footer):
            for x in part._element.iter(qn('w:tbl')):
                tw = x.find(qn('w:tblPr') + '/' + qn('w:tblW'))
                grid = sum(int(g.get(qn('w:w'))) for g in x.find(qn('w:tblGrid')).findall(qn('w:gridCol')))
                tcw = sum(int(e.get(qn('w:w'))) for e in x.iter(qn('w:tcW')))
                if tw is None or abs(int(tw.get(qn('w:w'))) - limit) > 3 or abs(grid - limit) > 3 or abs(tcw - limit) > 3: hf_bad.append((si + 1, tw.get(qn('w:w')) if tw is not None else None, grid, tcw, limit))
                elif sec.page_width > sec.page_height: land_cov += 1
    n_land = (dj or {}).get('landscape_sections') or 0
    check(f'{t}：docx 每节页眉页脚表 tblW / tblGrid / tcW = 该节版心宽（横向节 {n_land} 个，覆盖横向页眉页脚表 {land_cov} 张）', not hf_bad and (land_cov >= 2 * n_land), hf_bad[:3])
    import zipfile
    z = zipfile.ZipFile(path); allxml = ''.join(z.read(n).decode('utf-8', 'ignore') for n in z.namelist() if n.endswith('.xml')); z.close()
    check(f'{t}：docx 无 ⟪ ⟫ 残留，settings 有 updateFields，目录为 TOC 域', '⟪' not in allxml and '⟫' not in allxml and 'w:updateFields w:val="true"' in allxml and ' TOC \\o' in allxml)
    so = rb.get('soffice')
    if so is None:
        if soffice_required: missing_dep(f'{t}：docx LibreOffice 转换核验', 'LibreOffice（soffice）')
    else:
        check(f'{t}：docx 经 LibreOffice 转 PDF 成功，中文字体已嵌入（{", ".join((so.get("fonts") or [])[:3])}）', so.get('ok'), so)
    return d


def docx_preview(label, rd):
    if not DOCX_PREVIEWS: return
    pdfs = glob.glob(os.path.join(rd, 'out', '.work', 'docx-check', '*.pdf'))
    if pdfs:
        os.makedirs(DOCX_PREVIEWS, exist_ok=True)
        subprocess.run(['pdftoppm', '-r', '80', '-png', pdfs[0], os.path.join(DOCX_PREVIEWS, label)], check=False)


def find_sample(t):
    for root in (os.path.join(DOC_SHARED, 'types', t, 'samples'), os.path.join(STAGING, t, 'samples')):
        c = sorted(glob.glob(os.path.join(root, '*', 'doc.md')))
        if c: return os.path.dirname(c[0])
    return None


def choose_pack(t):
    real = os.path.join(DOC_SHARED, 'types', t)
    if os.path.exists(os.path.join(real, 'pack.json')):
        errs, _ = validate.validate_data('pack', json.load(open(os.path.join(real, 'pack.json'))), real)
        if not errs: return real, '正式类型包'
    return os.path.join(FIXTURE_PACKS, t), '自测类型包'


def clean_env():
    env = {k: v for k, v in os.environ.items() if k not in ('HTTP_PROXY', 'HTTPS_PROXY', 'http_proxy', 'https_proxy', 'ALL_PROXY', 'all_proxy')}
    nv = sorted(glob.glob(os.path.expanduser('~/.nvm/versions/node/v*/bin')), key=lambda x: [int(n) for n in re.findall(r'\d+', x)])
    if nv: env['PATH'] = nv[-1] + os.pathsep + env.get('PATH', '')
    return env


def page_texts(pdf):
    out = subprocess.run(['pdftotext', '-layout', pdf, '-'], capture_output=True, text=True).stdout.split('\f')
    return out[:-1] if out and not out[-1].strip() else out


def dark_bands(pdf, top_mm=17, bottom_mm=15, dpi=50):
    data = subprocess.run(['pdftoppm', '-f', '1', '-l', '1', '-r', str(dpi), '-gray', pdf], capture_output=True).stdout
    m = re.match(rb'P5\s+(\d+)\s+(\d+)\s+(\d+)\s', data)
    w, h = int(m.group(1)), int(m.group(2)); pix = data[m.end():]
    k = dpi / 25.4
    def dark(y0, y1): return sum(1 for y in range(y0, y1) for v in pix[y * w:(y + 1) * w] if v < 200)
    return dark(0, int(top_mm * k)), dark(h - int(bottom_mm * k), h)


def bookmark_count(r):
    def cnt(o): return sum(cnt(x) if isinstance(x, list) else 1 for x in o)
    return cnt(r.outline)


# ---------- 1. 解析器
print('== 解析器 ==')
pf = test_docmark_parse.run(verbose=True)
fails += ['docmark：' + x for x in pf]

# ---------- 2. 引擎业务词守卫（与 doc-shared 同一词表）
WORDS = ['报价', '确认单', '毛利', '底价', '人天', '售前', 'presales', 'pricing', '用例', 'reddit', '建站']
hits = []
for f in sorted(glob.glob(os.path.join(ROOT, 'scripts', '*'))):
    if not f.endswith(('.py', '.mjs')): continue
    for i, line in enumerate(open(f, encoding='utf-8').read().split('\n'), 1):
        hits += [f'{os.path.basename(f)}:{i} 「{w}」' for w in WORDS if w.lower() in line.lower()]
check('业务词守卫：doc-render/scripts 不出现业务词', not hits, hits[:8])

# ---------- 3. 不打印 PDF 的用例
print('== 渲染（不打印） ==')
PRD_PACK = os.path.join(FIXTURE_PACKS, 'prd')


def mkrun(name, files):
    d = os.path.join(BASE, name)
    for rel, content in files.items():
        p = os.path.join(d, rel); os.makedirs(os.path.dirname(p), exist_ok=True)
        open(p, 'w', encoding='utf-8').write(content if isinstance(content, str) else json.dumps(content, ensure_ascii=False))
    return d


DOCJ = {'type': 'prd', 'title': '自测', 'version': '1.0', 'status': 'draft', 'audience': 'internal', 'brand': 'internal', 'language': 'zh-CN',
        'owner': '自测', 'project': 'selftest', 'lark_folder': {'pending_reason': '自测', 'path': '自测'}}
small = svg('小字号', w=1180, px=12)
neg = mkrun('neg', {'doc.json': DOCJ, 'doc.md': '# 小字号\n\n## 图\n\n![小字号图](figures/small.svg){#fig:small}\n\n![未构建](figures/missing.mmd)\n\n见 @fig:small。\n\n| a | b |\n|---|---|\n| ⟪高亮⟫ | 未闭合⟪ |\n',
                    'figures/small.svg': small, 'figures/missing.mmd': 'flowchart TD\n  A-->B\n'})
c, o, e = run([sys.executable, RENDER, neg, '--printer', 'none', '--pack', PRD_PACK])
d = jout(o)
rj = json.load(open(os.path.join(neg, 'out', 'render.json'))) if os.path.exists(os.path.join(neg, 'out', 'render.json')) else {}
ly4 = [x for x in rj.get('layout_issues', []) if x['rule'] == 'LY4']
check('LY4 反例：12 px / 1180 宽 → 等效 5.25 pt 必改；未构建的图必改；退出码 3', c == 3 and len(ly4) == 2 and abs((rj['figures'][0].get('min_font_pt') or 0) - 5.25) < 0.06, (c, ly4, rj.get('figures'), o[-400:], e[-400:]))
check('printer none：pdf 为 null、目录页码为 null，render.json 过 schema', rj.get('pdf') is None and all(t['page'] is None for t in rj.get('toc', [])) and not validate.validate_data('render', rj)[0])
html_neg = open(os.path.join(neg, 'out', 'doc.html'), encoding='utf-8').read()
xml_neg = open(os.path.join(neg, 'out', 'feishu.xml'), encoding='utf-8').read()
check('非法高亮：HTML 与飞书 XML 都不泄漏 ⟪ ⟫', '⟪' not in html_neg + xml_neg and '⟫' not in html_neg + xml_neg)
check('SVG 内联：id 与 url(#) 加命名空间前缀', 'id="f1-arr"' in html_neg and 'url(#f1-arr)' in html_neg and 'class="f1-note"' in html_neg and '.f1-note{' in html_neg)
check('飞书 XML：.mmd 给 whiteboard type="mermaid" 源；.svg 给 type="svg"', 'type="mermaid" path="@./figures/missing.mmd"' in xml_neg and 'type="svg" path="@./figures/small.svg"' in xml_neg)
hl = json.load(open(os.path.join(neg, 'out', 'render.highlights.json'))) if os.path.exists(os.path.join(neg, 'out', 'render.highlights.json')) else rj.get('highlights', {})
check('highlights 统计：总数、每处长度', hl.get('total') == 1 and hl['items'][0]['length'] == 2, hl)

guard = mkrun('guard', {'doc.json': DOCJ, 'doc.md': '# 守卫\n\n<!-- include: internal-cost.md -->\n', 'internal-cost.md': '成本\n'})
c, o, e = run([sys.executable, RENDER, guard, '--printer', 'none', '--pack', PRD_PACK])
check('守卫：include 白名单外（internal）退出码 1 且不写 render.json', c == 1 and not os.path.exists(os.path.join(guard, 'out', 'render.json')), (c, o[-300:]))

st = mkrun('state', {'doc.json': DOCJ, 'doc.md': '# 状态\n\n## 一\n\n正文。\n'})
c1, o1, e1 = run([sys.executable, os.path.join(DOC_SHARED, 'scripts', 'run_state.py'), st, 'init', '--type', 'prd', '--mode', 'new', '--pack', PRD_PACK])
c, o, e = run([sys.executable, RENDER, st, '--printer', 'none', '--pack', PRD_PACK])
stage = json.load(open(os.path.join(st, 'run-state.json'))).get('stage') if os.path.exists(os.path.join(st, 'run-state.json')) else None
check('run-state：渲染完成后 stage=render', c1 == 0 and c == 0 and stage == 'render', (c1, o1[-300:], c, o[-300:], stage))

c, o, e = run([sys.executable, RENDER, os.path.join(BASE, 'nope')])
check('用法：运行目录不存在退出码 2', c == 2, o)

code_off = mkrun('code-off', {'doc.json': DOCJ, 'doc.md': '# 代码\n\n## 一\n\n```json\n{"a": 1}\n```\n'})
c, o, e = run([sys.executable, RENDER, code_off, '--printer', 'none', '--pack', PRD_PACK])
check('拒绝：类型包 code_blocks=false 时出现代码块，退出码 1 且不写产物', c == 1 and not os.path.exists(os.path.join(code_off, 'out', 'feishu.xml')), (c, o[-300:]))
esc_dir = mkrun('escape', {'doc.json': DOCJ, 'doc.md': '# 越界\n\n## 一\n\n![外部图片](../secret.png)\n'})
open(os.path.join(BASE, 'secret.png'), 'wb').write(bytes.fromhex('89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082'))
c, o, e = run([sys.executable, RENDER, esc_dir, '--printer', 'none', '--pack', PRD_PACK])
check('拒绝：图片路径越出运行目录，退出码 1，不内嵌', c == 1 and not os.path.exists(os.path.join(esc_dir, 'out', 'doc.html')), (c, o[-300:]))
png = bytes.fromhex('89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082')
misc = mkrun('misc', {'doc.json': DOCJ, 'doc.md': '# 杂项\n\n## 一\n\n<!-- widths: 1,4,1 -->\n| a | b | c |\n|---|---|---|\n| 1 | 2 | 3 |\n\n<!-- widths: 50,300,80 -->\n| x | y | z |\n|---|---|---|\n| 1 | 2 | 3 |\n\n<!-- data: data/m.csv -->\n\n<!-- grid: 2 -->\n![图一](shots/a.png)\n混入的段落文字\n<!-- /grid -->\n',
                      'data/m.csv': '列一,列二\n"第一行\n第二行",x\n'})
open(os.path.join(misc, 'shots', 'a.png') if os.path.isdir(os.path.join(misc, 'shots')) else (os.makedirs(os.path.join(misc, 'shots')) or os.path.join(misc, 'shots', 'a.png')), 'wb').write(png)
c, o, e = run([sys.executable, RENDER, misc, '--printer', 'none', '--pack', PRD_PACK])
xml = open(os.path.join(misc, 'out', 'feishu.xml'), encoding='utf-8').read() if c in (0, 3) else ''
hx = open(os.path.join(misc, 'out', 'doc.html'), encoding='utf-8').read() if c in (0, 3) else ''
cols = [c_ for c_ in (list(map(int, re.findall(r'<col width="(\d+)"', g))) for g in re.findall(r'<colgroup>(.*?)</colgroup>', xml)) if len(c_) == 3]  # 去掉封面文档控制表（两列）
check('飞书列宽：权重 1,4,1 与 50,300,80 都按权重归一化映射到约 820 px（不按数值猜像素）', len(cols) >= 2 and cols[0][1] > cols[0][0] * 2 and 780 <= sum(cols[0]) <= 860
      and 780 <= sum(cols[1]) <= 880 and cols[1][1] > cols[1][2] > cols[1][0] and cols[1] != [50, 300, 80], cols)
c, o, e = run([sys.executable, RENDER, misc, '--printer', 'none', '--pack', PRD_PACK, '--legacy-widths-px', '--no-state'])
xml_l = open(os.path.join(misc, 'out', 'feishu.xml'), encoding='utf-8').read() if c in (0, 3) else ''
cols_l = [c_ for c_ in (list(map(int, re.findall(r'<col width="(\d+)"', g))) for g in re.findall(r'<colgroup>(.*?)</colgroup>', xml_l)) if len(c_) == 3]
check('飞书列宽迁移开关 --legacy-widths-px：像素值 50,300,80 原样，权重 1,4,1 仍归一化', len(cols_l) >= 2 and cols_l[1] == [50, 300, 80] and 780 <= sum(cols_l[0]) <= 860, (c, cols_l))
check('CSV 单元格内换行：飞书 XML 与 HTML 都输出 <br/>', '第一行<br/>第二行' in xml and '第一行<br/>第二行' in hx)
check('grid 混入段落：HTML 与飞书 XML 都把段落放在 grid 之前、grid 只含图', hx.find('混入的段落文字') < hx.find('class="dm-grid"') and xml.find('混入的段落文字') < xml.find('<grid>') and '混入' not in re.search(r'<grid>.*?</grid>', xml, re.S).group(0), (c, o[-300:]))


# ---------- 3b. 1b 验收问题回归（第三波 W3-D1）
outside = os.path.join(BASE, 'outside'); os.makedirs(outside, exist_ok=True)
open(os.path.join(outside, 'secret.md'), 'w').write('# 机密正文\n\n## 一\n\n机密。\n')
open(os.path.join(outside, 'secret.png'), 'wb').write(png)
sl = mkrun('src-symlink', {'doc.json': DOCJ})
os.symlink(os.path.join(outside, 'secret.md'), os.path.join(sl, 'doc.md'))
c, o, e = run([sys.executable, RENDER, sl, '--printer', 'none', '--pack', PRD_PACK, '--no-state'])
check('拒绝：doc.md 是指向运行目录外的符号链接，退出码 1，不写任何产物', c == 1 and not os.path.exists(os.path.join(sl, 'out')), (c, o[-300:]))
fl = mkrun('fig-symlink', {'doc.json': DOCJ, 'doc.md': '# 图越界\n\n## 一\n\n![位图](shots/link.png)\n\n![矢量](figures/a.svg)\n\n![源](figures/b.mmd)\n',
                           'figures/b.mmd': 'flowchart TD\n  A-->B\n'})
os.makedirs(os.path.join(fl, 'shots'))
os.symlink(os.path.join(outside, 'secret.png'), os.path.join(fl, 'shots', 'link.png'))
open(os.path.join(outside, 'secret.svg'), 'w').write(svg('机密'))
os.symlink(os.path.join(outside, 'secret.svg'), os.path.join(fl, 'figures', 'a.svg'))
os.symlink(os.path.join(outside, 'secret.svg'), os.path.join(fl, 'figures', 'b.svg'))
c, o, e = run([sys.executable, RENDER, fl, '--printer', 'none', '--pack', PRD_PACK, '--no-state'])
fp = [x for x in jout(o).get('diagnostics', []) if x['code'] == 'figure-path']
check('拒绝：图文件（PNG、SVG、MMD 的 build 产物）经符号链接越界，各报 figure-path，退出码 1，不写飞书 XML', c == 1 and len(fp) == 3 and not os.path.exists(os.path.join(fl, 'out', 'feishu.xml')), (c, fp, o[-300:]))
# 飞书后端二次校验：解析通过之后图文件被换成外指链接
fs = mkrun('feishu-recheck', {'doc.json': DOCJ, 'doc.md': '# 二次校验\n\n## 一\n\n![位图](shots/a.png)\n', 'shots/.keep': ''})
open(os.path.join(fs, 'shots', 'a.png'), 'wb').write(png)
code = ("import sys, os; sys.path.insert(0, %r); import common, backend_feishu\n"
        "ctx = common.load_context(%r, %r); doc = common.dp.parse_file(ctx.run_dir, 'doc.md', pack=ctx.pack); common.finish_context(ctx, doc)\n"
        "p = os.path.join(ctx.run_dir, 'shots', 'a.png'); os.remove(p); os.symlink(%r, p)\n"
        "xml, st = backend_feishu.build_feishu(ctx); print('PATH' if 'shots/a.png\"' in xml else 'NOPATH', st['images'])") % (os.path.join(ROOT, 'scripts'), fs, PRD_PACK, os.path.join(outside, 'secret.png'))
c, o, e = run([sys.executable, '-c', code])
check('飞书后端二次校验：图文件被替换为外指链接时不写 path、不计图片', c == 0 and o.strip() == 'NOPATH 0', (c, o, e[-300:]))
# LY10：写死期望（受控别名精确匹配）
DOCJ_REV = {**DOCJ, 'revision_history': [{'version': '1.0', 'date': '2026-09-15', 'author': '自测', 'summary': '初稿'}]}
l10 = mkrun('ly10', {'doc.json': DOCJ_REV, 'doc.md': '# 重复页\n\n<!-- summary -->\n摘要内容。\n<!-- /summary -->\n\n## 摘要方法\n\n正文。\n\n## 文档控制策略\n\n正文。\n\n## 执行摘要\n\n正文。\n\n## 修订记录\n\n正文。\n\n## 文档控制与修订记录\n\n正文。\n\n#### 摘要\n\n正文。\n'})
c, o, e = run([sys.executable, RENDER, l10, '--printer', 'none', '--pack', PRD_PACK, '--no-state'])
r10 = json.load(open(os.path.join(l10, 'out', 'render.json'))) if os.path.exists(os.path.join(l10, 'out', 'render.json')) else {}
check('LY10：只报「执行摘要」「修订记录」「文档控制与修订记录」三处（h3 h4 h5）；「摘要方法」「文档控制策略」与三级标题「摘要」不报',
      sorted(x['target'] for x in r10.get('layout_issues', []) if x['rule'] == 'LY10') == ['h3', 'h4', 'h5'], [x for x in r10.get('layout_issues', []) if x['rule'] == 'LY10'])
th = mkrun('toc-hl', {'doc.json': DOCJ, 'doc.md': '# 目录高亮\n\n## ⟪重点章节⟫\n\n正文。\n\n## 二\n\n正文。\n'})
c, o, e = run([sys.executable, RENDER, th, '--printer', 'none', '--pack', PRD_PACK, '--no-state'])
rth = json.load(open(os.path.join(th, 'out', 'render.json'))) if os.path.exists(os.path.join(th, 'out', 'render.json')) else {}
check('高亮统计：标题里的一处高亮只计 1 次（目录副本不计数、不加定位标记）', (rth.get('highlights') or {}).get('total') == 1, rth.get('highlights'))

# 高亮颜色变体（2026-09-15）：! risk / + tip / ~ warn / ? decision，无前缀 neutral 与改动前完全一致
hk = mkrun('hl-kinds', {'doc.json': DOCJ, 'doc.md': '# 高亮颜色\n\n无⟪甲⟫，险⟪!乙⟫，利⟪+丙⟫，注⟪~丁⟫，决⟪?戊⟫。\n\n## ⟪+目录高亮⟫\n\n正文。\n\n## 二\n\n正文。\n'})
c, o, e = run([sys.executable, RENDER, hk, '--printer', 'none', '--docx', '--no-state', '--pack', PRD_PACK], timeout=600)
rhk = json.load(open(os.path.join(hk, 'out', 'render.json'))) if os.path.exists(os.path.join(hk, 'out', 'render.json')) else {}
hhk = open(os.path.join(hk, 'out', 'doc.html'), encoding='utf-8').read()
xhk = open(os.path.join(hk, 'out', 'feishu.xml'), encoding='utf-8').read()
check('高亮颜色：退出码 0，render.json 过 schema', c == 0 and rhk and not validate.validate_data('render', rhk)[0], (c, jout(o).get('must_fix') or jout(o).get('error'), e[-400:]))
# 计数减 1：CSS 规则里各出现一次
cls = {k: hhk.count(f'dm-hl-{k}') - 1 for k in ('risk', 'tip', 'warn', 'decision')}
check('高亮颜色 HTML：4 个前缀各出对应 class（tip 3 处 = 正文 + 标题 + 目录副本），无前缀仍是单独的 dm-hl',
      cls == {'risk': 1, 'tip': 3, 'warn': 1, 'decision': 1} and hhk.count('class="dm-hl"') == 1, (cls, hhk.count('class="dm-hl"')))
check('高亮颜色 HTML：底色与描边取 tokens callout 变量，未写死十六进制',
      all(f'var(--dm-callout-{k}-bg)' in hhk and f'var(--dm-callout-{k}-border)' in hhk for k in ('risk', 'tip', 'warn', 'decision'))
      and not re.search(r'mark\.dm-hl-\w+\{[^}]*#[0-9A-Fa-f]{6}', hhk))
fb = {m: len(re.findall(r'<span background-color="%s">' % m, xhk)) for m in ('light-purple', 'light-red', 'light-green', 'light-orange', 'light-blue')}
check('高亮颜色 飞书：5 种底色复用 callout 映射各出现（neutral=light-purple、risk=light-red、tip=light-green×2、warn=light-orange、decision=light-blue）',
      fb == {'light-purple': 1, 'light-red': 1, 'light-green': 2, 'light-orange': 1, 'light-blue': 1}, fb)
dhk = rhk.get('docx') or {}
dx8 = [x for x in (dhk.get('readback') or {}).get('checks', []) if x['rule'] == 'DX8' and '重点高亮' in x['name']]
check('高亮颜色 docx：DX8 统计不受 kind 影响（6 处合计），读回文本不含前缀字符',
      dhk.get('highlights') == 6 and dx8 and dx8[0]['ok'] and dx8[0]['actual'] == 6
      and (dhk.get('readback') or {}).get('highlights') == ['甲', '乙', '丙', '丁', '戊', '目录高亮'], (dhk.get('highlights'), dx8, (dhk.get('readback') or {}).get('highlights')))
check('高亮颜色：前缀字符不泄漏到任何输出', not any(ch in hhk + xhk for ch in '⟪⟫') and '⟪!' not in hhk and '>!乙' not in hhk and '>+丙' not in hhk)
code = ("import sys; sys.path.insert(0, %r); import common\n"
        "ctx = common.load_context(%r, %r); ctx.profile = dict(ctx.profile, bookmark_levels=2)\n"
        "doc = common.dp.parse_file(ctx.run_dir, 'doc.md', pack=ctx.pack)\n"
        "try:\n    common.finish_context(ctx, doc); print('NOERR')\nexcept common.RenderError as ex:\n    print('ERR', ex.code)") % (os.path.join(ROOT, 'scripts'), th, PRD_PACK)
c, o, e = run([sys.executable, '-c', code])
check('书签层级固定为 3：品牌档案写 bookmark_levels=2 时报用法错误（退出码 2）', o.strip() == 'ERR 2', (o, e[-300:]))
lt_tok = 'RULE_CONFLICT_EXCEEDED_LIMIT_ERROR_CODE'
fw = mkrun('feishu-wbr', {'doc.json': DOCJ, 'doc.md': f'# 长标识符\n\n## 一\n\n| 编号 | 错误码 |\n|---|---|\n| 1 | {lt_tok} |\n'})
c, o, e = run([sys.executable, RENDER, fw, '--printer', 'none', '--pack', PRD_PACK, '--no-state'])
xw = open(os.path.join(fw, 'out', 'feishu.xml'), encoding='utf-8').read() if c in (0, 3) else ''
hw = open(os.path.join(fw, 'out', 'doc.html'), encoding='utf-8').read() if c in (0, 3) else ''
check('飞书长标识符：原文完整、不含 <wbr>（HTML 端才加断行提示）', lt_tok in xw and '<wbr' not in xw and '<wbr>' in hw, (c, xw[-300:]))

# ---------- 3c. docx 后端（不打印 PDF）
NOPACK_DOCJ = {**DOCJ, 'type': 'docx-fixture', 'revision_history': [{'version': '1.0', 'date': '2026-09-15', 'author': '自测', 'summary': '初稿'}]}
feat_md = ('# docx 全块自测\n\n<!-- summary -->\n一句话摘要，含⟪重点结论⟫。\n<!-- /summary -->\n\n## 概述 {#sec:intro}\n\n'
           '正文见 @fig:flow 与 @tbl:cmp，脚注一[^n1]，再次引用[^n1]，外链 [官网](https://www.example.com)。\n\n'
           '需求 R1 在此定义{#req:R1}。实体引用 @req:R1、@req:R2，四级标题引用 @sec:deep。\n\n'
           '1. 第一步\n   - 子项甲\n   - 子项乙\n2. 第二步\n\n- 无序一\n- 无序二\n\n'
           '> [!warn] 注意事项\n> - 列表在高亮块里\n\n```json\n{"a": 1, "b": "<!-- 代码里的注释不算残留 -->"}\n```\n\n'
           '![流程图](figures/flow.mmd){#fig:flow}\n\n<!-- table: 对比 {#tbl:cmp} -->\n| 项 | 说明 |\n|---|---|\n| ⟪A⟫ | 甲 |\n| B{#req:R2} | 乙 |\n\n'
           '<!-- pagebreak -->\n\n## 横向 {#sec:wide}\n\n<!-- landscape -->\n\n### 宽表\n\n| c1 | c2 | c3 | c4 | c5 | c6 | c7 | c8 |\n|---|---|---|---|---|---|---|---|\n| TC-A-01 | 2 | 3 | 4 | 5 | 6 | 7 | 8 |\n\n#### 子节\n\n三级标题正文。\n\n##### 细节 {#sec:deep}\n\n四级标题正文（DocMark ##### 为四级）。\n\n'
           '<!-- /landscape -->\n\n## 网格\n\n<!-- grid: 2 -->\n![图一](shots/a.png)\n![图二](shots/b.png)\n![图三](shots/c.png)\n<!-- /grid -->\n\n## 附录 A：术语\n\n术语表。\n\n[^n1]: 脚注正文，含 [链接](https://www.example.com)。\n')
fx = mkrun('docx-feat', {'doc.json': NOPACK_DOCJ, 'doc.md': feat_md, 'figures/flow.mmd': 'flowchart TD\n  A-->B\n'})
build_svgs(fx)
os.makedirs(os.path.join(fx, 'shots'), exist_ok=True)
for nm in 'abc': open(os.path.join(fx, 'shots', f'{nm}.png'), 'wb').write(png_bytes(320, 200))
c, o, e = run([sys.executable, RENDER, fx, '--printer', 'none', '--docx', '--no-state', '--keep-work'], timeout=600)
rfx = json.load(open(os.path.join(fx, 'out', 'render.json'))) if os.path.exists(os.path.join(fx, 'out', 'render.json')) else {}
dfx = rfx.get('docx') or {}
check('docx 全块：退出码 0，render.json 过 schema', c == 0 and rfx and not validate.validate_data('render', rfx)[0], (c, jout(o).get('must_fix') or jout(o).get('error'), e[-400:]))
check('docx 全块：原生脚注 1 个 + NOTEREF 再次引用 1 次、横向节 1、代码块 1、高亮块 1、目录域、图 4 张（含 grid 3 张两行）、高亮 2 处（摘要与表格）',
      dfx.get('footnotes') == 1 and dfx.get('footnote_repeats') == 1 and dfx.get('landscape_sections') == 1 and dfx.get('code_blocks') == 1 and dfx.get('callouts') == 1
      and dfx.get('toc') and dfx.get('images') == 4 and dfx.get('highlights') == 2, {k: v for k, v in dfx.items() if k != 'readback'})
fx_doc = __import__('docmark_parse').parse_file(fx, 'doc.md')
docx_checks('docx 全块', fx, rfx, fx_doc, expect_all_passed=False,
            meta_expect={'Title': 'docx 全块自测', 'Author': NOPACK_DOCJ['owner'], 'Subject': NOPACK_DOCJ['type'], 'Keywords': f"{NOPACK_DOCJ['project']}, {NOPACK_DOCJ['type']}, v{NOPACK_DOCJ['version']}"})
hfx = open(os.path.join(fx, 'out', 'doc.html'), encoding='utf-8').read() if os.path.exists(os.path.join(fx, 'out', 'doc.html')) else ''
check('HTML 实体交叉引用：正文锚点处发射 id（ent-req-R1、ent-req-R2 各一次），引用链接指向它们，四级标题引用指向 sec-deep',
      hfx.count('id="ent-req-R1"') == 1 and hfx.count('id="ent-req-R2"') == 1 and 'href="#ent-req-R1"' in hfx and 'href="#ent-req-R2"' in hfx and 'href="#sec-deep"' in hfx and 'id="sec-deep"' in hfx)
rbx = dfx.get('readback') or {}
blank_chk = [c_ for c_ in rbx.get('checks', []) if c_['rule'] == 'DX8' and '空白页' in c_['name']]
ref_chk = [c_ for c_ in rbx.get('checks', []) if c_['rule'] == 'DX8' and c_['name'].startswith('交叉引用：已解析引用')]
sox = rbx.get('soffice') or {}
check('docx readback：LibreOffice 空白页检查项存在且与 soffice.blank_pages 一致（DX8 建议）；all_passed 与各检查项一致；交叉引用逐条检查通过',
      ((not sox) or (blank_chk and blank_chk[0]['severity'] == '建议' and blank_chk[0]['ok'] == (not sox.get('blank_pages'))))
      and rbx.get('all_passed') == all(c_['ok'] for c_ in rbx.get('checks', []))
      and ref_chk and ref_chk[0]['ok'] and ref_chk[0]['expected'] == f"{sum(1 for r_ in fx_doc.refs if r_['resolved'])} 条", (blank_chk, ref_chk))
import zipfile as _zf
if dfx.get('path') and os.path.exists(os.path.join(fx, dfx['path'])):
    zz = _zf.ZipFile(os.path.join(fx, dfx['path'])); dxml = zz.read('word/document.xml').decode(); fxml = zz.read('word/footnotes.xml').decode() if 'word/footnotes.xml' in zz.namelist() else ''
    rels = zz.read('word/_rels/document.xml.rels').decode(); ctx_ = zz.read('[Content_Types].xml').decode(); zz.close()
    ext_ids = re.findall(r'Id="(rId\d+)"[^>]*Target="https://www\.example\.com"[^>]*TargetMode="External"|TargetMode="External"[^>]*Target="https://www\.example\.com"[^>]*Id="(rId\d+)"', rels)
    ext_ids = {a_ or b_ for a_, b_ in ext_ids} | set(re.findall(r'Id="(rId\d+)"[^>]*Type="[^"]*/hyperlink"[^>]*Target="https://www\.example\.com"', rels))
    bms = set(re.findall(r'<w:bookmarkStart w:id="\d+" w:name="([^"]+)"', dxml))
    links_ = set(re.findall(r'w:anchor="([^"]+)"', dxml))
    hidden_targets = {'_fig_flow', '_tbl_cmp', '_ent_req_R1', '_ent_req_R2', '_sec_deep'}
    check('docx 全块：分页符、横向 orient；图、表、实体锚点、四级标题的交叉引用链接都指向存在的隐藏书签（_ 开头）', '<w:pageBreakBefore/>' in dxml and 'w:orient="landscape"' in dxml
          and hidden_targets <= links_ and hidden_targets <= bms, (sorted(hidden_targets - bms), sorted(hidden_targets - links_)))
    n_le3 = sum(1 for h_ in fx_doc.headings if h_['level'] <= 3)
    visible_ = {b_ for b_ in bms if not b_.startswith('_')}
    check(f'docx 全块：可见书签数 {len(visible_)} = 一至三级标题数 {n_le3}，四级标题不是可见书签', len(visible_) == n_le3 and 'sec_deep' not in visible_ and any(h_['level'] == 4 for h_ in fx_doc.headings), sorted(visible_))
    check('docx 全块：外链在 document.xml.rels 中为 External 关系且正文 hyperlink 引用该 rId', bool(ext_ids) and any(f'r:id="{i}"' in dxml for i in ext_ids), rels[:400])
    check('docx 全块：脚注部件有文档关系与 Content-Type，引用 id=1 对应脚注正文，再次引用为指向书签 _Ref_dmfn1 的 NOTEREF 域',
          'Target="footnotes.xml"' in rels and 'PartName="/word/footnotes.xml"' in ctx_ and '<w:footnoteReference w:id="1"/>' in dxml
          and '<w:footnote w:id="1">' in fxml and '脚注正文' in fxml and ' NOTEREF _Ref_dmfn1 ' in dxml and '_Ref_dmfn1' in bms)
    so = (dfx.get('readback') or {}).get('soffice') or {}
    check('docx 全块：LibreOffice 转出 PDF 含横向页', not so or so.get('landscape_pages', 0) >= 1, so)
    dx8_blank = [x for x in rfx.get('layout_issues', []) if x['rule'] == 'DX8' and '空白页' in x['message']]
    check(f'docx 全块：layout_issues 的 DX8 空白页记录与 blank_pages={so.get("blank_pages")} 一致（有则建议级、无则不报）', not so or (bool(so.get('blank_pages')) == bool(dx8_blank) and all(x['severity'] == '建议' for x in dx8_blank)), (so.get('blank_pages'), dx8_blank))
docx_preview('fixture-all-blocks', fx)
mut_code = r'''
import sys, os, re, zipfile, json
sys.path.insert(0, %r)
import common, svgtools, render, backend_docx
from docx import Document
ctx = common.load_context(%r); doc = common.dp.parse_file(ctx.run_dir, 'doc.md', pack=ctx.pack); common.finish_context(ctx, doc); render.prepare_figures(ctx, common, svgtools)
out = os.path.join(%r, 'mut'); os.makedirs(out, exist_ok=True)
base = os.path.join(out, 'base.docx'); stats = backend_docx.build_docx(ctx, base)
def failed(path):
    rb, _ = backend_docx.readback(ctx, path, stats)
    return sorted({c['rule'] + ':' + c['name'] for c in rb['checks'] if not c['ok']}), rb['ok'], rb['all_passed']
def edit(name, fn):
    p = os.path.join(out, name + '.docx'); zi = zipfile.ZipFile(base); zo = zipfile.ZipFile(p, 'w', zipfile.ZIP_DEFLATED)
    for it in zi.infolist():
        b = zi.read(it.filename)
        zo.writestr(it, fn(it.filename, b.decode('utf-8')).encode('utf-8') if it.filename.endswith('.xml') else b)
    zo.close(); return p
res = {'base': failed(base)}
res['no_entity_bookmark'] = failed(edit('nobm', lambda n, x: re.sub(r'<w:bookmarkStart w:id="\d+" w:name="_ent_req_R1"/>', '', x) if n == 'word/document.xml' else x))
d = Document(base); d.core_properties.author = '别人'; p = os.path.join(out, 'author.docx'); d.save(p); res['author'] = failed(p)
land = next(s for s in Document(base).sections if s.page_width > s.page_height)
hdr = str(land.header.part.partname).lstrip('/')
res['landscape_header_width'] = failed(edit('hdrw', lambda n, x: re.sub(r'(<w:tblW w:type="dxa" w:w=")(\d+)', lambda m: m.group(1) + str(int(m.group(2)) - 2000), x, count=1) if n == hdr else x))
res['visible_h4'] = failed(edit('h4', lambda n, x: x.replace('w:name="_sec_deep"', 'w:name="sec_deep"').replace('w:anchor="_sec_deep"', 'w:anchor="sec_deep"') if n == 'word/document.xml' else x))
print(json.dumps(res, ensure_ascii=False))
''' % (os.path.join(ROOT, 'scripts'), fx, BASE)
c, o, e = run([sys.executable, '-c', mut_code], timeout=300)
mres = jout(o)
def has_fail(key, rule, word): return any(x.startswith(rule) and word in x for x in (mres.get(key) or [[]])[0])
check('docx 反例（篡改后读回）：删掉实体锚点书签 → DX8 交叉引用失败；改 Author → DX7 必改；横向节页眉表宽 -2000 → DX5 必改；四级标题书签改为可见 → DX1 必改',
      bool(mres.get('base')) and not mres['base'][0] and has_fail('no_entity_bookmark', 'DX8', '交叉引用')
      and has_fail('author', 'DX7', '') and mres['author'][1] is False and has_fail('landscape_header_width', 'DX5', '') and has_fail('visible_h4', 'DX1', '可见书签'), (c, o[-800:], e[-400:]))
unit_code = r'''
import sys, types, json
sys.path.insert(0, %r)
import backend_docx as bd
ctx = types.SimpleNamespace(profile={'header': {'left': '{logo}', 'right': '{website}'}, 'footer': {'left': '第 {page} 页，共 {pages} 页', 'right': '- {page} - · {brand}'}}, vals={'website': 'www.x.com', 'brand': '易点 Brand'})
latin_fonts = 'name type encoding emb sub uni object ID\n---- ---- ---- --- --- --- ------\nBAAAAA+LiberationSans TrueType WinAnsi yes yes yes 13 0\n'
cjk_noemb = latin_fonts + 'CAAAAA+PingFangSC-Regular CID TrueType Identity-H no no yes 14 0\n'
cjk_emb = latin_fonts + 'DAAAAA+PingFangSC-Regular CID TrueType Identity-H yes yes yes 15 0\n'
print(json.dumps({
  'blank_custom': bd.page_residue(ctx, 'www.x.com\n第 6 页，共 9 页\n- 6 - · 易点 Brand\n'),
  'body_custom': bd.page_residue(ctx, 'www.x.com\n正文一句\n第 6 页，共 9 页\n'),
  'latin_ok': bd.font_problems('Hello world 12', latin_fonts)[0],
  'cjk_missing': bd.font_problems('中文正文', latin_fonts)[0],
  'cjk_not_embedded': bd.font_problems('中文正文', cjk_noemb)[0],
  'cjk_ok': bd.font_problems('中文正文', cjk_emb)[0],
  'no_pdffonts': bd.font_problems('Hello', None)[0]}, ensure_ascii=False))
''' % os.path.join(ROOT, 'scripts')
c, o, e = run([sys.executable, '-c', unit_code])
ur = jout(o)
merge_code = r'''
import sys, json
sys.path.insert(0, %r)
import backend_docx as bd
base = lambda: {'checks': [{'rule': 'DX1', 'name': 'x', 'ok': True, 'severity': '必改'}, {'rule': 'DX8', 'name': 'y', 'ok': True, 'severity': '建议'}]}
out = {'blank': bd.merge_soffice(base(), {'ok': True, 'blank_pages': [6]}), 'convert_fail': bd.merge_soffice(base(), {'ok': False, 'error': 'x', 'blank_pages': []}),
       'none': bd.merge_soffice(base(), None)}
print(json.dumps({k: {'ok': v['ok'], 'all_passed': v['all_passed'], 'checks': [(c['rule'], c['ok'], c['severity']) for c in v['checks']]} for k, v in out.items()}, ensure_ascii=False))
''' % os.path.join(ROOT, 'scripts')
c, o, e = run([sys.executable, '-c', merge_code])
mg = jout(o)
check('readback 合并 LibreOffice 结果：空白页 → DX8 建议项失败，ok 仍为 True、all_passed 为 False；转换失败 → DX9 必改，ok 为 False；未安装 → 不加检查项',
      mg.get('blank', {}).get('ok') is True and mg['blank']['all_passed'] is False and ['DX8', False, '建议'] in mg['blank']['checks']
      and mg.get('convert_fail', {}).get('ok') is False and ['DX9', False, '必改'] in mg['convert_fail']['checks']
      and mg.get('none', {}).get('ok') is True and mg['none']['all_passed'] is True and len(mg['none']['checks']) == 2, (o, e[-300:]))
check('空白页判定兼容自定义页脚模板（「第 {page} 页，共 {pages} 页」「- {page} -」）：只剩页眉页脚为空白，有正文不算', ur.get('blank_custom') == '' and ur.get('body_custom') == '正文一句', (o, e[-300:]))
check('字体核验：纯拉丁文 PDF 不要求中文字体；含中文时只有嵌入（emb=yes）的中文字体才通过；缺 pdffonts 判失败',
      ur.get('latin_ok') == [] and bool(ur.get('cjk_missing')) and bool(ur.get('cjk_not_embedded')) and ur.get('cjk_ok') == [] and bool(ur.get('no_pdffonts')), ur)
mp = mkrun('docx-missing-png', {'doc.json': NOPACK_DOCJ, 'doc.md': '# 缺图\n\n## 一\n\n![未出 PNG](figures/x.svg)\n', 'figures/x.svg': svg('x')})
c, o, e = run([sys.executable, RENDER, mp, '--printer', 'none', '--docx', '--no-state'], timeout=300)
rmp = json.load(open(os.path.join(mp, 'out', 'render.json'))) if os.path.exists(os.path.join(mp, 'out', 'render.json')) else {}
check('docx 反例：doc-figures 未产出 PNG → DX4 必改、退出码 3、docx 放占位不中断', c == 3 and any(x['rule'] == 'DX4' and x['severity'] == '必改' for x in rmp.get('layout_issues', [])) and (rmp.get('docx') or {}).get('images') == 0, (c, rmp.get('layout_issues')))
code = ("import sys, types; sys.path.insert(0, %r); import render\n"
        "A = lambda v: types.SimpleNamespace(docx=v); C = lambda pack, prof: types.SimpleNamespace(pack=pack, profile=prof)\n"
        "r = [render.want_docx(A(None), C({'outputs': []}, {'outputs': ['docx']})), render.want_docx(A(None), C({}, {'outputs': ['docx']})),\n"
        "     render.want_docx(A(None), C(None, {})), render.want_docx(A(True), C({'outputs': []}, {})), render.want_docx(A(False), C({'outputs': ['docx']}, {'outputs': ['docx']}))]\n"
        "print(r)") % os.path.join(ROOT, 'scripts')
c, o, e = run([sys.executable, '-c', code])
check('docx 默认输出优先级：类型包 outputs=[] 覆盖品牌档案；类型包无字段时回落品牌档案；都没有为否；--docx / --no-docx 最高', o.strip() == '[False, True, False, True, False]', (o, e[-300:]))
pk_out = mkrun('pack-outputs', {'pack.json': {**json.load(open(os.path.join(PRD_PACK, 'pack.json'))), 'outputs': ['docx']}, 'qa-rules.md': 'x\n'})
errs_pk, _ = validate.validate_data('pack', json.load(open(os.path.join(pk_out, 'pack.json'))), pk_out)
check('pack.schema：outputs 字段只收 docx，写 ["docx"] 不引入新的 schema 错误', not [x for x in errs_pk if 'outputs' in str(x)], errs_pk[:3])
od = mkrun('docx-default', {'doc.json': DOCJ, 'doc.md': '# 默认输出\n\n## 一\n\n正文。\n'})
c, o, e = run([sys.executable, RENDER, od, '--printer', 'none', '--pack', pk_out, '--no-state'], timeout=300)
c2, o2, e2 = run([sys.executable, RENDER, od, '--printer', 'none', '--pack', pk_out, '--no-state', '--no-docx'], timeout=300)
od_path = os.path.join(od, jout(o).get('docx') or '-')
check('docx 默认输出：类型包 outputs 含 docx 时不加参数也输出（退出码 0、文件存在、读回 ok）；--no-docx 覆盖为不输出',
      c == 0 and bool(jout(o).get('docx')) and os.path.isfile(od_path) and jout(o).get('docx_readback') is True
      and jout(o2).get('docx') is None and json.load(open(os.path.join(od, 'out', 'render.json')))['docx'] is None, (c, jout(o).get('docx'), jout(o).get('docx_readback'), jout(o2).get('docx')))

# ---------- 4. 三份样张出 HTML + PDF + XML
if QUICK:
    skip('样张 PDF 渲染', '--quick')
elif not (shutil.which('pdftotext') and shutil.which('pdftoppm')):
    missing_dep('样张 PDF 渲染', 'poppler（pdftotext、pdftoppm）')
else:
    print('== 样张渲染（PDF） ==')
    lark = shutil.which('lark-cli', path=clean_env()['PATH'])
    for t in SAMPLE_TYPES:
        src = find_sample(t)
        if not src:
            check(f'{t}：找到样张', False, '样张目录不存在'); continue
        rd = os.path.join(BASE, t, os.path.basename(src))
        shutil.copytree(src, rd, ignore=shutil.ignore_patterns('out'))
        build_svgs(rd)
        n_hl = 0
        if t == 'prd':
            p = os.path.join(rd, 'doc.md'); s = open(p, encoding='utf-8').read()
            lines = s.split('\n'); k = next(i for i, l in enumerate(lines) if l.startswith('## '))
            lines.insert(k + 1, '\n本期两个核心指标：⟪积分月活占比回到 35%⟫，⟪兑换失败率低于 2%⟫。\n')
            open(p, 'w', encoding='utf-8').write('\n'.join(lines)); n_hl = 2
        pack, which = choose_pack(t)
        c, o, e = run([sys.executable, RENDER, rd, '--pack', pack, '--keep-work', '--docx'], timeout=900)
        d = jout(o)
        check(f'{t}：渲染完成且无版式必改（{which}）', c == 0 and d.get('ok'), (c, d.get('must_fix') or d.get('error'), e[-300:]))
        rjp = os.path.join(rd, 'out', 'render.json')
        if not os.path.exists(rjp):
            continue
        rj = json.load(open(rjp))
        errs, _ = validate.validate_data('render', rj)
        check(f'{t}：render.json 通过 schema', not errs, errs[:3])
        html = open(os.path.join(rd, 'out', 'doc.html'), encoding='utf-8').read()
        xml = open(os.path.join(rd, 'out', 'feishu.xml'), encoding='utf-8').read()
        tabs_h = re.findall(r'<table\b[^>]*>(.*?)</table>', html, re.S); tabs_x = re.findall(r'<table\b[^>]*>(.*?)</table>', xml, re.S)
        check(f'{t}：HTML {len(tabs_h)} 个表格、飞书 XML {len(tabs_x)} 个表格全部有 thead', tabs_h and tabs_x and all('<thead' in x for x in tabs_h + tabs_x))
        pdfp = os.path.join(rd, rj['pdf']['path'])
        r = PdfReader(pdfp)
        heads = [x for x in json.load(open(os.path.join(rd, 'out', 'render.json')))['toc']]
        from importlib import import_module
        sys.path.insert(0, os.path.join(DOC_SHARED, 'scripts'))
        import docmark_parse as dp
        doc = dp.parse_file(rd, 'doc.md', pack=json.load(open(os.path.join(pack, 'pack.json'))))
        expect = sum(1 for h in doc.headings if h['level'] <= 3)
        check(f'{t}：合并后书签数 {bookmark_count(r)} = 一至三级标题数 {expect}（layout.md §4）', bookmark_count(r) == expect == rj['pdf']['bookmarks'])
        check(f'{t}：封面页没有残留链接注释（/Annots）', r.pages[0].get('/Annots') is None)
        work = os.path.join(rd, 'out', '.work')
        passes = sorted(glob.glob(os.path.join(work, 'pass*.pdf')), key=lambda x: int(re.findall(r'pass(\d+)', x)[0]))
        if passes:
            fr = PdfReader(os.path.join(work, 'full.pdf')); lr = PdfReader(passes[-1])
            flat = []
            def walk(items):
                for it in items:
                    if isinstance(it, list): walk(it)
                    else: flat.append(fr.get_destination_page_number(it) + 1)
            walk(fr.outline)
            toc_by_title = [x['page'] for x in rj['toc']]
            heads_all = [h for h in doc.headings]
            full_pages_for_toc = [flat[i] for i, h in enumerate(heads_all) if h['level'] <= 2] if len(flat) == len(heads_all) else None
            check(f'{t}：定位遍（带标记）与终遍（无标记）页数一致、标题页一致', len(lr.pages) == len(fr.pages) and full_pages_for_toc == toc_by_title, (len(lr.pages), len(fr.pages), full_pages_for_toc, toc_by_title))
        check(f'{t}：页数与 render.json 一致（{len(r.pages)}）', len(r.pages) == rj['pdf']['pages'])
        md = r.metadata or {}
        check(f'{t}：PDF 元数据 Title、Author、Subject', md.get('/Title') and md.get('/Author') and md.get('/Subject'), dict(md))
        texts = page_texts(pdfp)
        bad = [(x['title'], x['page']) for x in rj['toc'] if not x['page'] or norm(x['title'])[:30] not in norm(texts[x['page'] - 1])]
        toc_page = next((i + 1 for i, tx in enumerate(texts) if norm(tx).startswith('目录') or '目录' in norm(tx)[:60]), None)
        toc_nums = norm(texts[toc_page - 1]) if toc_page else ''
        check(f'{t}：目录 {len(rj["toc"])} 项都有页码，且 pdftotext 在该页找到标题', rj['toc'] and not bad, bad[:5])
        check(f'{t}：目录页码都在目录页（第 {toc_page} 页）之后', toc_page and all(x['page'] > toc_page for x in rj['toc']))
        check(f'{t}：目录页印出了页码（目录页文字含每项「标题…页码」）', toc_page and all(norm(x['title'])[:12] in toc_nums for x in rj['toc']) and all(str(x['page']) in toc_nums for x in rj['toc']))
        first = texts[0]
        foot2 = [l for l in texts[1].split('\n') if re.search(r'\d+\s*/\s*%d\s*$' % len(r.pages), l)]
        foot_text = norm(re.sub(r'\d+\s*/\s*\d+\s*$', '', foot2[-1])) if foot2 else ''
        check(f'{t}：封面 pdftotext 不含第 2 页页脚的整行文字', foot_text and foot_text not in norm(first), (foot_text, first[-200:]))
        check(f'{t}：封面无「1 / N」页码且第 2 页有页脚页码', not re.search(r'(^|\n)\s*1\s*/\s*\d+\s*($|\n)', first) and re.search(r'2\s*/\s*%d' % len(r.pages), texts[1]), texts[1][-120:])
        top, bottom = dark_bands(pdfp)
        check(f'{t}：封面页眉区与页脚区像素无印迹（暗像素 顶 {top} / 底 {bottom}）', top == 0 and bottom == 0)
        check(f'{t}：输出无 ⟪ ⟫ 残留（HTML、XML、PDF 文字）', not any(ch in html + xml + ''.join(texts) for ch in '⟪⟫'))
        check(f'{t}：render.json 无版式必改', not [x for x in rj['layout_issues'] if x['severity'] == '必改'], rj['layout_issues'][:3])
        hlp = os.path.join(rd, 'out', 'render.highlights.json')
        hl = rj.get('highlights') or (json.load(open(hlp)) if os.path.exists(hlp) else {})
        check(f'{t}：highlights 统计总数 {hl.get("total")} = 插入数 {n_hl}，且每处有页码', hl.get('total') == n_hl and all(x['page'] for x in hl.get('items', [])))
        check(f'{t}：docx 重点高亮 {(rj.get("docx") or {}).get("highlights")} 处 = 插入数 {n_hl}', (rj.get('docx') or {}).get('highlights') == n_hl)
        ALIASES = {'修订记录', '文档控制', '文档控制与修订记录', '摘要', '执行摘要'}   # 写死的受控别名（layout.md LY10），不复用引擎算法
        dup_heads = [h for h in doc.headings if h['level'] <= 2 and norm(h['plain']) in ALIASES]
        ly10 = [x for x in rj['layout_issues'] if x['rule'] == 'LY10']
        check(f'{t}：LY10 只报与受控别名整串相同的一二级标题（{len(dup_heads)} 个），逐个带页码', len(ly10) == len(dup_heads) and all(isinstance(x['page'], int) for x in ly10), ly10)
        docx_checks(t, rd, rj, doc)
        docx_preview(t, rd)
        check(f'{t}：一级标题分页模式为 auto（technical 封面默认）', d.get('h1_page_break') == 'auto', d.get('h1_page_break'))
        if t == 'tech-spec':
            check('tech-spec：代码块进飞书 XML（pre lang + code）', '<pre lang="json"><code>' in xml and rj['feishu']['code_blocks'] == 1)
            check('tech-spec：docx 代码块 1 个（DM Code 段落）', (rj.get('docx') or {}).get('code_blocks') == 1, rj.get('docx'))
            check('tech-spec：5 个画板，其中 Mermaid 3 个', rj['feishu']['whiteboards'] == 5 and rj['feishu']['whiteboards_mermaid'] == 3, rj['feishu'])
            fig_nums = [b_['number'] for b_ in doc.figures]   # 期望取解析器（上游）编号：样张章节由类型包维护方调整过
            check('tech-spec：图编号「图 章-序」与等效字号', fig_nums and all(re.fullmatch(r'图 \d+-\d+', x) for x in fig_nums) and [f['number'] for f in rj['figures']] == fig_nums and all(f['min_font_pt'] and f['min_font_pt'] >= 7 for f in rj['figures']), rj['figures'][:2])
            check(f'tech-spec：交叉引用渲染为「如{fig_nums[0] if fig_nums else ""} 所示」', fig_nums and f'如 <a class="dm-ref" href="#fig-current-arch">{fig_nums[0]}</a> 所示' in html)
        if t == 'test-cases':
            land = [i + 1 for i, p in enumerate(r.pages) if float(p.mediabox.width) > float(p.mediabox.height)]
            check(f'test-cases：横向页段生效（横向页 {len(land)} 页）', len(land) >= 5)
            check('test-cases：10 列宽表 dense 且 landscape', any(x['columns'] == 10 and x['dense'] and x['landscape'] for x in rj['tables']))
            dso = ((rj.get('docx') or {}).get('readback') or {}).get('soffice') or {}
            check(f'test-cases：docx 横向节 1 个，LibreOffice 转出横向页 {dso.get("landscape_pages")} 页', (rj.get('docx') or {}).get('landscape_sections') == 1 and (not dso or dso.get('landscape_pages', 0) >= 5), rj.get('docx'))
        if lark:
            c2, o2, e2 = run([lark, 'docs', '+script', '--command', 'parse', '--content', '@./out/feishu.xml', '--format', 'json'], cwd=rd, env=clean_env(), timeout=120)
            dd = jout(o2)
            status = ((dd.get('data') or {}).get('assessment') or {}).get('status')
            check(f'{t}：飞书 XML 过 lark-cli docs +script parse（本地，不发布）', c2 == 0 and status == 'passed', (c2, status, (dd.get('data') or {}).get('diagnostics'), e2[-200:]))
        else:
            missing_dep(f'{t}：lark-cli parse', 'lark-cli')
        if PREVIEWS:
            os.makedirs(PREVIEWS, exist_ok=True)
            for label, pg in (('cover', 1), ('toc', toc_page or 2)):
                outp = os.path.join(PREVIEWS, f'{t}-{label}')
                subprocess.run(['pdftoppm', '-f', str(pg), '-l', str(pg), '-r', '110', '-png', '-singlefile', pdfp, outp], check=False)
            for label, pg in (('body', min(len(r.pages), (toc_page or 2) + 1 + (5 if t == 'test-cases' else 3))),):
                subprocess.run(['pdftoppm', '-f', str(pg), '-l', str(pg), '-r', '110', '-png', '-singlefile', pdfp, os.path.join(PREVIEWS, f'{t}-{label}')], check=False)

    # tech-spec 用 always 再渲染一次：auto 页数更少、稀疏页更少
    ts = find_sample('tech-spec')
    if ts:
        pack_ts, _ = choose_pack('tech-spec')
        pk = json.load(open(os.path.join(pack_ts, 'pack.json'))); pk['features']['h1_page_break'] = 'always'
        pdir = mkrun('pack-always', {'pack.json': pk, 'qa-rules.md': 'x\n'})
        rd_a = os.path.join(BASE, 'always', 'ts'); shutil.copytree(ts, rd_a, ignore=shutil.ignore_patterns('out')); build_svgs(rd_a)
        c, o, e = run([sys.executable, RENDER, rd_a, '--pack', pdir, '--no-state'], timeout=900)
        ra = json.load(open(os.path.join(rd_a, 'out', 'render.json'))) if os.path.exists(os.path.join(rd_a, 'out', 'render.json')) else {}
        ru = json.load(open(os.path.join(BASE, 'tech-spec', os.path.basename(ts), 'out', 'render.json'))) if os.path.exists(os.path.join(BASE, 'tech-spec', os.path.basename(ts), 'out', 'render.json')) else {}
        n9a = sum(1 for x in ra.get('layout_issues', []) if x['rule'] == 'LY9'); n9u = sum(1 for x in ru.get('layout_issues', []) if x['rule'] == 'LY9')
        check(f'分页：always 模式报出稀疏页 LY9（{n9a} 页），auto 模式页数更少（{ru.get("pdf", {}).get("pages")} < {ra.get("pdf", {}).get("pages")}）且稀疏页更少（{n9u}）',
              ra and ru and n9a >= 1 and ru['pdf']['pages'] < ra['pdf']['pages'] and n9u < n9a, (c, o[-300:]))


    # 分页模式与 LY9（1b 验收 #7 #8）：断言每个一级标题的位置，期望按规则写死，不复用 pdf.py 的判定
    def pack_with(mode):
        pk = json.load(open(os.path.join(PRD_PACK, 'pack.json'))); pk['features'] = {**pk['features'], 'h1_page_break': mode}
        return mkrun(f'pack-{mode}', {'pack.json': pk, 'qa-rules.md': 'x\n'})
    para = '会员积分规则引擎把获取比例、抵扣比例、上限与有效期从代码中抽出，运营在后台配置后即时生效，灰度与定时生效另行说明。' * 3
    counts = [7, 1, 9, 1, 5, 1, 12, 1, 3, 2]
    md = '# 分页位置\n\n' + ''.join(f'## 第{k + 1}章\n\n' + ''.join(f'{para}\n\n' for _ in range(n)) for k, n in enumerate(counts))
    def last_pass_marks(rdir):
        ps = sorted(glob.glob(os.path.join(rdir, 'out', '.work', 'pass*.pdf')), key=lambda x: int(re.findall(r'pass(\d+)', x)[0]))
        if not ps: return {}, []
        bb = subprocess.run(['pdftotext', '-bbox', ps[-1], '-'], capture_output=True, text=True).stdout
        marks, hs = {}, []
        for pi, m in enumerate(re.finditer(r'<page width="([\d.]+)" height="([\d.]+)">(.*?)</page>', bb, re.S)):
            hs.append(float(m.group(2)))
            for y0, t_ in re.findall(r'<word xMin="[\d.]+" yMin="([\d.]+)" xMax="[\d.]+" yMax="[\d.]+">(.*?)</word>', m.group(3)):
                for key in re.findall(r'QQ([a-z]\d+)QQ', t_): marks.setdefault(key, (pi + 1, float(y0)))
        return marks, hs
    pos = {}
    for mode in ('always', 'never', 'auto'):
        rdm = mkrun(f'break-{mode}', {'doc.json': DOCJ, 'doc.md': md})
        c, o, e = run([sys.executable, RENDER, rdm, '--pack', pack_with(mode), '--no-state', '--keep-work', '--no-feishu'], timeout=900)
        rjm = json.load(open(os.path.join(rdm, 'out', 'render.json'))) if os.path.exists(os.path.join(rdm, 'out', 'render.json')) else {}
        marks, hs = last_pass_marks(rdm)
        pos[mode] = [(marks.get(f'e{i}'), marks.get(f'h{i}'), hs) for i in range(1, len(counts))]
        check(f'分页 {mode}：渲染完成且无 LY1', rjm.get('pdf') and not [x for x in rjm.get('layout_issues', []) if x['rule'] == 'LY1'] and all(p_[0] and p_[1] for p_ in pos[mode]), (c, jout(o).get('must_fix'), pos[mode][:2]))
    gap = 9 * 72 / 25.4
    ok_always = all(h_[0] == e_[0] + 1 for e_, h_, _ in pos['always'] if e_ and h_)
    ok_never = all(h_[0] == e_[0] or (h_[0] == e_[0] + 1 and e_[1] > hs_[e_[0] - 1] * 0.8) for e_, h_, hs_ in pos['never'] if e_ and h_) and any(h_ and e_ and h_[0] == e_[0] for e_, h_, _ in pos['never'])
    bad_auto = []
    for k, (e_, h_, hs_) in enumerate(pos['auto']):
        if not (e_ and h_): continue
        low = e_[1] + gap > hs_[e_[0] - 1] * 0.65
        if low and h_[0] != e_[0] + 1: bad_auto.append((k + 2, '应换页', e_, h_))
        if not low and h_[0] != e_[0]: bad_auto.append((k + 2, '不应换页（上一章在本页 65% 以上结束）', e_, h_))
    check('分页 always：第 2 章起每个一级标题都在上一章结束页的下一页', ok_always, pos['always'])
    check('分页 never：一级标题都紧接上一章（至少一处同页），不强制换页', ok_never, pos['never'])
    check('分页 auto：上一章结束在 65% 以下的才换页，结束在 65% 以上的不换页（固定点收敛，break 可增可减）', not bad_auto, bad_auto)
    pbd = mkrun('ly9-pagebreak', {'doc.json': DOCJ, 'doc.md': '# 手工分页\n\n## A\n\n短短一句。\n\n<!-- pagebreak -->\n\n## B\n\n' + ''.join(f'{para}\n\n' for _ in range(6)) + '## C\n\n' + ''.join(f'{para}\n\n' for _ in range(6))})
    c, o, e = run([sys.executable, RENDER, pbd, '--pack', pack_with('never'), '--no-state', '--no-feishu'], timeout=900)
    rpb = json.load(open(os.path.join(pbd, 'out', 'render.json'))) if os.path.exists(os.path.join(pbd, 'out', 'render.json')) else {}
    pa = next((x['page'] for x in rpb.get('toc', []) if x['title'] == 'A'), None)
    check(f'LY9：never 模式下以手工分页结束的稀疏页（A 所在第 {pa} 页）照报，不当自然章末豁免', pa and any(x['rule'] == 'LY9' and x['page'] == pa for x in rpb.get('layout_issues', [])), rpb.get('layout_issues'))
    cellx = '规则参数从代码中抽出为独立服务，账户与兑换服务同步查询并带本地缓存兜底，兑换扣减用数据库行锁保证原子性。' * 4
    tblx = '| 项 | 说明 |\n|---|---|\n' + ''.join(f'| 第{k}项 | {cellx} |\n' for k in range(8))
    sp9 = mkrun('ly9-samepage', {'doc.json': DOCJ, 'doc.md': '# 稀疏页\n\n## A\n\n短短一句。\n\n' + tblx + '\n## B\n\n短句。\n\n' + ''.join(f'{para}\n\n' for _ in range(3)) + '## C\n\n' + ''.join(f'{para}\n\n' for _ in range(8))})
    c, o, e = run([sys.executable, RENDER, sp9, '--pack', pack_with('never'), '--no-state', '--keep-work', '--no-feishu'], timeout=900)
    r9 = json.load(open(os.path.join(sp9, 'out', 'render.json'))) if os.path.exists(os.path.join(sp9, 'out', 'render.json')) else {}
    m9, _ = last_pass_marks(sp9)
    h0p = (m9.get('h0') or (None,))[0]; t0p = (m9.get('t0') or (None,))[0]; e1p = (m9.get('e1') or (None,))[0]; h1p = (m9.get('h1') or (None,))[0]
    check(f'LY9 反例：A 章跨两页（标题第 {h0p} 页、表格第 {t0p} 页）、B 标题与 A 末段同页（e1={e1p}，h1={h1p}）→ A 的稀疏首页第 {h0p} 页照报 LY9',
          bool(h0p) and t0p == e1p == h1p == h0p + 1 and any(x['rule'] == 'LY9' and x['page'] == h0p for x in r9.get('layout_issues', [])), (c, m9, [x for x in r9.get('layout_issues', []) if x['rule'] == 'LY9']))

    token = 'RULE_CONFLICT_EXCEEDED_LIMIT_ERROR_CODE'
    lt = mkrun('longtoken', {'doc.json': DOCJ, 'doc.md': f'# 长标识符\n\n## 一\n\n<!-- widths: 1,1,4 -->\n| 编号 | 错误码 | 说明 |\n|---|---|---|\n| 1 | {token} | 说明文字 |\n| 2 | ABCDEFGHIJKLMNOPQRSTUVWXYZABCDEFGHIJKLMN | 无分隔的超长串 |\n\n### 二级\n\n#### 三级\n\n正文。\n\n## 二\n\n### 二之一\n\n正文。\n'})
    c, o, e = run([sys.executable, RENDER, lt, '--pack', PRD_PACK, '--no-state'], timeout=600)
    rl = json.load(open(os.path.join(lt, 'out', 'render.json'))) if os.path.exists(os.path.join(lt, 'out', 'render.json')) else {}
    if rl.get('pdf'):
        pdfp = os.path.join(lt, rl['pdf']['path'])
        bb = subprocess.run(['pdftotext', '-bbox', pdfp, '-'], capture_output=True, text=True).stdout
        words = [w for w in re.findall(r'>([^<>]+)</word>', bb) if len(w) >= 2 and w in token]
        pieces = []
        for w in words:
            if ''.join(pieces + [w]) == token[:len(''.join(pieces + [w]))]: pieces.append(w)
            if ''.join(pieces) == token: break
        check('长标识符：只在 _ 之后断行，不在字母中间断开', ''.join(pieces) == token and len(pieces) >= 2 and all(p_.endswith('_') for p_ in pieces[:-1]), pieces)
        check('长标识符：无分隔的超长串强制断行，不溢出版心（无 LY7）', not [x for x in rl['layout_issues'] if x['rule'] == 'LY7'], rl['layout_issues'])
        tree = []
        rr = PdfReader(pdfp)
        def walk2(items, dep):
            for it in items:
                if isinstance(it, list): walk2(it, dep + 1)
                else: tree.append((dep, it.title, rr.get_destination_page_number(it) + 1))
        walk2(rr.outline, 1)
        tx = page_texts(pdfp)
        check('书签树：层级 1/2/3/1/2，目标页上能找到对应标题，无 LY2', [x[0] for x in tree] == [1, 2, 3, 1, 2] and all(norm(ti.split(' ', 1)[-1]) in norm(tx[pg - 1]) for _, ti, pg in tree) and not [x for x in rl['layout_issues'] if x['rule'] == 'LY2'], (tree, rl['layout_issues']))
    else:
        check('长标识符与书签树用例渲染成功', False, (c, o[-400:], e[-300:]))

    # 旧文档（marketing 封面 + cyberklick 档案）：golden smoke-site 输入
    tgz = os.path.join(GOLDEN, 'inputs', 'smoke-site.tgz')
    if os.path.exists(tgz):
        sd = os.path.join(BASE, 'legacy'); os.makedirs(sd)
        subprocess.run(['tar', 'xzf', tgz, '-C', sd], check=True)
        run_dir = os.path.dirname(next(iter(glob.glob(os.path.join(sd, '**', 'proposal.md'), recursive=True))))
        pk = mkrun('legacy-pack', {'pack.json': {**json.load(open(os.path.join(PRD_PACK, 'pack.json'))), 'id': 'legacy', 'source_file': 'proposal.md', 'meta_file': 'brief.json',
                                                  'cover': 'marketing', 'brand_profile': 'cyberklick', 'include_allow': ['sections/*.md', 'pricing/*.md', 'terms.md', 'confirm-list.md'],
                                                  'features': {**json.load(open(os.path.join(PRD_PACK, 'pack.json')))['features'], 'h1_new_page': False}}, 'qa-rules.md': 'x\n'})
        c, o, e = run([sys.executable, RENDER, run_dir, '--pack', pk, '--profile', 'cyberklick', '--legacy-tilde', '--legacy-widths-px', '--no-state', '--docx', '--keep-work'], timeout=900)
        d = jout(o)
        rj = json.load(open(os.path.join(run_dir, 'out', 'render.json'))) if os.path.exists(os.path.join(run_dir, 'out', 'render.json')) else {}
        codes = {x['rule'] for x in rj.get('layout_issues', []) if x['severity'] == '必改'}
        check('旧文档（marketing 封面）：零改动渲染出 PDF，无 LY1 / LY3 / LY6 必改', c in (0, 3) and rj.get('pdf') and not ({'LY1', 'LY3', 'LY6'} & codes), (c, d.get('must_fix') or d.get('error')))
        resolved_old = open(os.path.join(GOLDEN, 'smoke-site', 'proposal.resolved.md'), encoding='utf-8').read()
        resolved_new = open(os.path.join(run_dir, 'out', 'proposal.resolved.md'), encoding='utf-8').read() if rj else ''
        check('旧文档：proposal.resolved.md 与 golden 逐字一致', resolved_old == resolved_new)
        if rj:
            import docmark_parse as _dp
            dl = docx_checks('旧文档 smoke-site（marketing 封面）', run_dir, rj, _dp.parse_file(run_dir, 'proposal.md', pack=json.load(open(os.path.join(pk, 'pack.json')))))
            if dl is not None:
                from docx.oxml.ns import qn as _qn
                ctl = [x for x in dl.element.body.iter(_qn('w:tbl')) if (x.find(_qn('w:tblPr') + '/' + _qn('w:tblStyle')) is not None and x.find(_qn('w:tblPr') + '/' + _qn('w:tblStyle')).get(_qn('w:val')) == {s_.name: s_.style_id for s_ in dl.styles}.get('DM Control'))]
                check('旧文档 docx：marketing 封面无文档控制表，封面有官网文字', not ctl and 'www.cyberklick.com' in ''.join(t_.text or '' for t_ in dl.element.body.iter(_qn('w:t'))))
            docx_preview('golden-smoke-site-marketing', run_dir)
        if PREVIEWS and rj.get('pdf'):
            subprocess.run(['pdftoppm', '-f', '1', '-l', '1', '-r', '110', '-png', '-singlefile', os.path.join(run_dir, rj['pdf']['path']), os.path.join(PREVIEWS, 'legacy-marketing-cover')], check=False)
    else:
        skip('旧文档 marketing 封面', 'golden inputs 不存在')

# ---------- 4d. PDF 视觉回归（W3-G：scripts/visual_regress.py；--update-visual-baseline 用第 4 节三份样张 PDF 重建 golden/visual 基线）
VISUAL = os.path.join(ROOT, 'scripts', 'visual_regress.py')
VISUAL_GOLDEN = os.path.join(GOLDEN, 'visual')
UPDATE_VISUAL = '--update-visual-baseline' in args


def vr(*a):
    c, o, e = run([sys.executable, VISUAL] + [str(x) for x in a], timeout=600)
    d = jout(o)
    if not d: d = {'raw': (o + e)[-400:]}
    return c, d


if QUICK:
    skip('视觉回归', '--quick')
elif not shutil.which('pdftoppm'):
    missing_dep('视觉回归', 'poppler（pdftoppm）')
else:
    print('== 视觉回归 ==')
    vpdf = {}
    for t in SAMPLE_TYPES:
        src = find_sample(t)
        rjp = os.path.join(BASE, t, os.path.basename(src), 'out', 'render.json') if src else ''
        if rjp and os.path.exists(rjp) and (json.load(open(rjp)).get('pdf') or {}).get('path'):
            vpdf[t] = os.path.join(BASE, t, os.path.basename(src), json.load(open(rjp))['pdf']['path'])
    check('视觉回归：第 4 节三份样张都渲染出 PDF', set(vpdf) == set(SAMPLE_TYPES), sorted(vpdf))
    for t, pdf in sorted(vpdf.items()):
        bdir = os.path.join(VISUAL_GOLDEN, t)
        if UPDATE_VISUAL:
            c, d = vr(pdf, '--baseline', bdir, '--update-baseline')
            check(f'{t}：golden 视觉基线已更新（{d.get("pages")} 页，{bdir}）', c == 0 and d.get('updated'), d)
        elif not os.path.exists(os.path.join(bdir, 'baseline.json')):
            missing_dep(f'{t}：与 golden 视觉基线比对', f'基线不存在 {bdir}（run_tests.py --update-visual-baseline 建立）')
        else:
            c, d = vr(pdf, '--baseline', bdir, '--out', os.path.join(BASE, 'visual-diff', t))
            worst = max((p.get('diff_ratio') or 0 for p in d.get('pages', [])), default=None)
            check(f'{t}：与 golden 视觉基线逐页比对通过（{(d.get("page_count") or {}).get("pdf")} 页，最大差异比例 {worst}）', c == 0 and d.get('ok'),
                  (c, d.get('error'), d.get('page_count'), d.get('failed_pages'), d.get('heatmap_dir')))
    ts_pdf = vpdf.get('tech-spec')
    if ts_pdf:
        vb = os.path.join(BASE, 'visual', 'self')
        c, d = vr(ts_pdf, '--baseline', vb, '--update-baseline')
        meta = json.load(open(os.path.join(vb, 'baseline.json'))) if c == 0 else {}
        check('视觉回归：--update-baseline 写入逐页 PNG 与 baseline.json（dpi、页数、栅格化环境）', c == 0 and meta.get('pages') == d.get('pages') and meta.get('env', {}).get('pillow')
              and len(glob.glob(os.path.join(vb, 'page-*.png'))) == meta.get('pages'), (c, d))
        c, d = vr(ts_pdf, '--baseline', vb, '--out', os.path.join(BASE, 'visual', 'diff-self'))
        check('视觉回归：同一份 PDF 与自身基线比对，每页差异像素为 0、退出 0、不写热图', c == 0 and d.get('ok') and all(p['status'] == 'same' and p['diff_pixels'] == 0 for p in d.get('pages', []))
              and not glob.glob(os.path.join(BASE, 'visual', 'diff-self', '*.png')), (c, d.get('failed_pages'), d.get('error')))
        c, d = vr(ts_pdf, '--baseline', vb, '--dpi', '72', '--out', os.path.join(BASE, 'visual', 'diff-dpi'))
        check('视觉回归：--dpi 与基线不同给 warning 并按基线 dpi 比较', c == 0 and any('dpi' in w for w in d.get('warnings', [])), d.get('warnings'))
        # 人为改文字：同宽数字替换（不重排）→ 检出差异但不超 0.5% 阈值；插入整段（重排）→ 判失败
        ts_src = find_sample('tech-spec'); pack_ts, _ = choose_pack('tech-spec')
        def rerender(label, fn):
            rdv = os.path.join(BASE, 'visual', label, 'ts'); shutil.copytree(ts_src, rdv, ignore=shutil.ignore_patterns('out')); build_svgs(rdv)
            p = os.path.join(rdv, 'doc.md'); s0 = open(p, encoding='utf-8').read(); s1 = fn(s0); assert s1 != s0
            open(p, 'w', encoding='utf-8').write(s1)
            c0, o0, e0 = run([sys.executable, RENDER, rdv, '--pack', pack_ts, '--no-state'], timeout=900)
            rj0 = json.load(open(os.path.join(rdv, 'out', 'render.json'))) if os.path.exists(os.path.join(rdv, 'out', 'render.json')) else {}
            return os.path.join(rdv, rj0['pdf']['path']) if (rj0.get('pdf') or {}).get('path') else None
        one = rerender('one-char', lambda s: s.replace('平均周期 9 个工作日', '平均周期 8 个工作日', 1))
        if one:
            c, d = vr(one, '--baseline', vb, '--out', os.path.join(BASE, 'visual', 'diff-one'))
            changed = [p for p in d.get('pages', []) if p['diff_pixels']]
            check(f'视觉回归：改一个字（9→8，不重排）检出差异：{[(p["page"], p["diff_pixels"], p["status"]) for p in changed]}，有热图，未超 0.5% 阈值不判失败',
                  c == 0 and len(changed) >= 1 and all(p['status'] == 'diff' and p['heatmap'] and os.path.exists(p['heatmap']) and p['bbox'] for p in changed)
                  and not d.get('page_count', {}).get('changed'), (c, d.get('error'), changed[:3]))
        else:
            check('视觉回归：改一个字的副本渲染出 PDF', False)
        para = '\n\n本段为视觉回归自测插入的整段文字，用来验证正文重排会被判为失败。' * 6 + '\n'
        big = rerender('reflow', lambda s: s.replace('\n## 2. ', para + '\n## 2. ', 1))
        if big:
            c, d = vr(big, '--baseline', vb, '--out', os.path.join(BASE, 'visual', 'diff-reflow'))
            check(f'视觉回归：插入整段文字（重排）判失败，退出 3，失败页 {d.get("failed_pages")}', c == 3 and not d.get('ok') and d.get('failed_pages'), (c, d.get('error'), d.get('failed_pages')))
            c, d = vr(ts_pdf, '--baseline', vb, '--out', os.path.join(BASE, 'visual', 'diff-reflow'))
            check('视觉回归：再次比较前清掉旧热图（相同 PDF 比对后热图目录无 PNG）', c == 0 and not glob.glob(os.path.join(BASE, 'visual', 'diff-reflow', '*.png')))
        else:
            check('视觉回归：插入整段的副本渲染出 PDF', False)
        # 页数不同：删掉最后一页
        from pypdf import PdfWriter
        cut = os.path.join(BASE, 'visual', 'cut.pdf'); w = PdfWriter(); rd_ = PdfReader(ts_pdf)
        for pg in rd_.pages[:-1]: w.add_page(pg)
        with open(cut, 'wb') as f: w.write(f)
        c, d = vr(cut, '--baseline', vb, '--out', os.path.join(BASE, 'visual', 'diff-cut'))
        last = (d.get('pages') or [{}])[-1]
        check('视觉回归：页数不同（少最后一页）→ page_count.changed、末页 removed 且失败、退出 3', c == 3 and d.get('page_count', {}).get('changed') and last.get('status') == 'removed' and last.get('page') in d.get('failed_pages', []),
              (c, d.get('page_count'), last))
        c, d = vr(ts_pdf, '--baseline', os.path.join(BASE, 'visual', 'cutbase'), '--update-baseline')
        c2, d2 = vr(cut, '--baseline', os.path.join(BASE, 'visual', 'cutbase'), '--update-baseline')
        n_png = len(glob.glob(os.path.join(BASE, 'visual', 'cutbase', 'page-*.png')))
        c3, d3 = vr(ts_pdf, '--baseline', os.path.join(BASE, 'visual', 'cutbase'), '--out', os.path.join(BASE, 'visual', 'diff-added'))
        check('视觉回归：--update-baseline 删除多余旧页；多一页时该页 added 且失败', c == 0 and c2 == 0 and d2.get('removed_files') and n_png == d2.get('pages')
              and c3 == 3 and (d3.get('pages') or [{}])[-1].get('status') == 'added', (d2, n_png, c3, (d3.get('pages') or [{}])[-1]))
        broken = os.path.join(BASE, 'visual', 'broken'); shutil.copytree(vb, broken); os.remove(os.path.join(broken, 'page-002.png'))
        c, d = vr(ts_pdf, '--baseline', broken)
        c2, d2 = vr(ts_pdf, '--baseline', os.path.join(BASE, 'visual', 'nowhere'))
        open(os.path.join(broken, 'baseline.json'), 'w').write('{bad')
        c3, d3 = vr(ts_pdf, '--baseline', broken)
        check('视觉回归：基线缺页、基线不存在、baseline.json 损坏 → 退出 2', (c, c2, c3) == (2, 2, 2) and '缺页' in d.get('error', '') and '不存在' in d2.get('error', ''), (c, d, c2, c3))
        c, d = vr(ts_pdf, '--baseline', vb, '--threshold', '1.5')
        check('视觉回归：参数越界退出 2', c == 2, d)
        from PIL import Image
        sized = os.path.join(BASE, 'visual', 'sized'); shutil.copytree(vb, sized)
        with Image.open(os.path.join(sized, 'page-001.png')) as im0:
            canvas = Image.new('RGB', (im0.size[0], im0.size[1] + 20), (255, 255, 255)); canvas.paste(im0.convert('RGB'), (0, 0))
        canvas.save(os.path.join(sized, 'page-001.png'))
        c, d = vr(ts_pdf, '--baseline', sized, '--out', os.path.join(BASE, 'visual', 'diff-sized'))
        p1 = (d.get('pages') or [{}])[0]
        check('视觉回归：页面尺寸变化而新增区域全白（差异像素 0）→ 仍判失败、有热图、退出 3', c == 3 and p1.get('size_changed') and p1.get('status') == 'fail' and p1.get('heatmap') and os.path.exists(p1['heatmap']), (c, p1))
    else:
        check('视觉回归：tech-spec 样张 PDF 可用于检出验证', False)

print(f"\n{'ALL PASS' if not fails else str(len(fails)) + ' FAILED'}" + (f'（SKIP {len(skips)}）' if skips else '') + f'（临时目录：{BASE}）')
if not KEEP and not fails:
    shutil.rmtree(BASE, ignore_errors=True)
sys.exit(1 if fails else 0)
